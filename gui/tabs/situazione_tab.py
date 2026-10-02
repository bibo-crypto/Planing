"""
situazione_tab.py
Dyeing Situation Tracker — embedded as the "Situazione" tab of the Delta
Dyeing PDF/Excel converter. Replaces the old Old Situazione / WOORKSHEET /
New Situazione Excel + Power Query chain: each source is uploaded and
validated separately, and per-Partita Old/New Comment history is kept in
a local SQLite database (see situazione_db.py) instead of copy-pasted
sheets.
"""
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinter import font as tkfont
from datetime import datetime
from typing import Callable
import pandas as pd

import utility.situazione_db as db
import parsers.situazione_loaders as data_loaders
import calculate.situazione as business_logic
import calculate.reports as reports
from calculate.abbina_suggestions import build_suggestions
from gui.tabs.yarn_shortage_tab import YarnShortageTab
from utility.utils import keep_window_on_top, logger
from utility.excel_io import safe_save_workbook
from .situazione_sources import SOURCE_BUTTON_NAMES, SOURCE_ORDER, SituationSourcesMixin
from .situazione_refresh import SituationRefreshMixin

STATUS_COLORS = {
    "Filato": "#e0e0e0",
    "pronto da spedire": "#c6efce",
    "Spedita": "#bdd7ee",
    "spedita": "#bdd7ee",
    "Check": "#ffc7ce",
    "C.Q": "#fff2cc",
}


def color_for_status(status):
    if not status:
        return "#ffffff"
    if status.startswith("Ritinta"):
        return "#d9c6f0"
    return STATUS_COLORS.get(status, "#ffffff")


# (internal_key, header shown in the app / exported file, cell type)
# Order and headers match the real "New Situazione" sheet exactly.
COLUMN_SPEC = [
    ("cliente", "CLIENTE", "text"),
    ("articolo", "Articolo", "text"),
    ("titolo", "Titolo", "text"),
    ("codice", "Codice", "number"),
    ("colore", "Colore", "text"),
    ("prezzo", "Prezzo Ord.", "text"),
    ("prezzo_lisini", "Prezzo Listini", "text"),
    ("densita", "Densita` (360-390)", "number"),
    ("ordine", "Ordine", "text"),
    ("riga", "Riga", "number"),
    ("data", "Data", "date"),
    ("delivery_date", "Delivery Date", "date"),
    ("consegna", "Consegna", "date"),
    ("partita", "Partita", "number"),
    ("rocche", "Rocche", "number"),
    ("mc", "M/C", "number"),
    ("comment", "Comment", "text"),
    ("raw_yarn_match", "Filato Disponibile", "text"),
    ("cq", "C.Q", "text"),
    ("tinto", "Tinto", "date"),
    ("bagno", "Bagno", "text"),
    ("old_comment", "Old Comm.", "text"),
    ("new_comment", "New Comm.", "text"),
    ("planedate", "PlaneDate", "text"),
    ("data_qualita", "Data Qualita", "date"),
    ("data_uscita", "Data Uscita", "date"),
    ("custom", "Custom", "text"),
    ("days_in_qc", "Days in Q.C", "number"),
    ("ritardo_consegna", "Ritardo (gg)", "number"),
]
COLUMN_LABELS = {key: header for key, header, _ in COLUMN_SPEC}
COLUMN_TYPES = {key: ctype for key, _, ctype in COLUMN_SPEC}


class SourceRow(ttk.Frame):
    """Compact upload card: file name button + current upload status."""

    def __init__(self, master, key, label, on_upload):
        super().__init__(master)
        self.key = key
        self.on_upload = on_upload

        self.status_var = tk.StringVar(value="Not uploaded")
        self.dot = tk.Label(self, text="●", fg="#9e9e9e", bg="#f4f6f8", font=("Segoe UI", 11))
        self.dot.grid(row=1, column=0, padx=(4, 2), sticky="w")

        button_name = SOURCE_BUTTON_NAMES.get(key, label)
        self.button = ttk.Button(
            self,
            text=button_name,
            command=self._browse,
            width=6,
            min_width=88,
            font=("Segoe UI", 8, "bold"),
            height=30,
        )
        self.button.grid(row=0, column=0, columnspan=2, padx=4, pady=(3, 2), sticky="ew")

        self.status_lbl = ttk.Label(self, textvariable=self.status_var, anchor="center",
                        foreground="#667085", width=8, font=("Segoe UI", 8))
        self.status_lbl.grid(row=1, column=1, padx=(2, 4), sticky="w")

        self.columnconfigure(1, weight=1)

    def _browse(self):
        path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls")])
        if not path:
            return
        self.on_upload(self.key, path)

    def set_status(self, ok, message):
        self.dot.config(fg="#2e7d32" if ok else "#c62828")
        self.status_var.set(message)


class SituazioneTab(SituationSourcesMixin, SituationRefreshMixin, ttk.Frame):
    """Embeddable 'Situazione' tab — hosted inside gui.py's main Notebook."""

    def __init__(self, master, on_shared_cache_changed: Callable[[], None] | None = None, on_notification=None):
        super().__init__(master)

        self._configure_styles()

        self._on_shared_cache_changed = on_shared_cache_changed
        self._on_notification = on_notification
        # Set post-construction from gui.py once MagazinoFilatoTab exists
        # (it's built after this tab). Used only to auto-fill the "Filato
        # Disponibile" column -- read lazily, so it's fine if it's not set
        # yet the first time _load_table_from_db() runs.
        self.magazino_tab = None

        db.init_db()

        self.loaded_frames = {}   # key -> DataFrame (validated, ready to merge)
        self.sort_state = {}      # column -> ascending bool
        self.current_df = pd.DataFrame()
        self._filter_after_id = None
        self._shared_syncing = False
        self._other_shared_syncing = False
        self._shared_source_paths = {}
        self._startup_restore_in_progress = False
        self._sources_restore_started = False
        self._startup_snapshot_current = False
        self._table_loaded_callbacks = []
        self._data_revision = 0
        self._tree_render_generation = 0
        self._tree_render_after_id = None
        self._shared_dfm_path = ""
        self._shared_prod_path = ""
        self._copertura_revision = 0
        cached_copertura = db.load_frame_cache("copertura")
        if isinstance(cached_copertura, pd.DataFrame) and not cached_copertura.empty:
            self.loaded_frames["copertura"] = cached_copertura
            self._copertura_revision += 1
        self._child_windows = {}
        self._price_densita_syncing = False

        self._build_upload_panel()
        self._build_toolbar()
        self._build_treeview()
        self._refresh_source_labels_from_db()
        # Show the last SQLite snapshot immediately.  Excel files are restored
        # later in a worker, so the application opens against the cache first
        # and refreshes when the latest source files have finished loading.
        self._load_table_from_db()

    def on_shown(self) -> None:
        """Restore shared Excel sources when the Situation page is opened."""
        if self._sources_restore_started:
            return
        self._sources_restore_started = True
        self.after_idle(self._auto_restore_saved_files)
        # Stagger the optional background syncs: starting all Excel readers at
        # the same instant made the UI and disk compete during startup.
        self.after(500, self.sync_shared_async)
        self.after(1200, self.sync_remaining_shared_sources)
        self.after(2200, self._recompute_prezzo_densita_async)
        self.after(3200, self.refresh_raw_yarn_match_async)

    # ------------------------------------------------------------------ UI
    def _configure_styles(self):
        # Note: the app's main window (ui/gui.py) already sets the global
        # ttk theme to "clam" once at startup for the same reason (the
        # native Windows theme ignores Treeview heading background colors).
        # Calling it again here is redundant but harmless (idempotent) —
        # kept as a safety net in case this tab is ever instantiated on its
        # own, outside the main app.
        style = ttk.Style(self)
        if "clam" in style.theme_names():
            style.theme_use("clam")
        style.configure("Upload.TLabelframe", background="#f4f6f8", borderwidth=1, relief="solid")
        style.configure("Upload.TLabelframe.Label", background="#f4f6f8", foreground="#344054",
                        font=("Segoe UI", 10, "bold"))
        style.configure("Situazione.Treeview", background="#ffffff", fieldbackground="#ffffff",
                        foreground="#1d2939", rowheight=29, borderwidth=1, relief="solid",
                        bordercolor="#b8c6d6", lightcolor="#b8c6d6", darkcolor="#b8c6d6",
                        font=("Segoe UI", 9))
        style.configure("Situazione.Treeview.Heading", background="#16324f", foreground="#ffffff",
                        relief="raised", borderwidth=1, padding=(8, 8),
                        font=("Segoe UI", 9, "bold"))
        style.map("Situazione.Treeview.Heading",
                  foreground=[("pressed", "#ffffff"), ("active", "#ffffff"), ("!active", "#ffffff")],
                  background=[("pressed", "#0b2239"), ("active", "#244b70"), ("!active", "#16324f")])
        style.map("Situazione.Treeview", background=[("selected", "#2563eb")],
                  foreground=[("selected", "#ffffff")])

    def _build_upload_panel(self):
        upload_area = ttk.Frame(self)
        upload_area.pack(side="top", fill="x", padx=8, pady=(2, 0))
        upload_area.columnconfigure(0, weight=3, uniform="upload_panels")
        upload_area.columnconfigure(1, weight=1, uniform="upload_panels")

        panel = ttk.LabelFrame(upload_area, text="1) Required files", style="Upload.TLabelframe")
        panel.grid(row=0, column=0, sticky="nsew")
        panel.configure(padding=(5, 1))

        self.source_rows = {}
        for index, key in enumerate(SOURCE_ORDER):
            label, _ = data_loaders.LOADERS[key]
            row = SourceRow(panel, key, label, self._handle_upload)
            row.grid(row=0, column=index, padx=2, pady=1, sticky="ew")
            self.source_rows[key] = row
        for column in range(len(SOURCE_ORDER)):
            panel.columnconfigure(column, weight=1, uniform="required_sources")

        codes_panel = ttk.LabelFrame(upload_area, text="Optional file", style="Upload.TLabelframe")
        codes_panel.grid(row=0, column=1, sticky="nsew", padx=(8, 0))
        codes_panel.configure(padding=(5, 1))
        codes_panel.columnconfigure(0, weight=1, uniform="optional_sources")
        codes_panel.columnconfigure(1, weight=1, uniform="optional_sources")

        codes_row = SourceRow(codes_panel, "codes", "Yarn codes (Articoli) - optional",
                               self._handle_codes_upload)
        codes_row.grid(row=0, column=0, padx=3, pady=1, sticky="ew")
        self.codes_row = codes_row

        listini_row = SourceRow(codes_panel, "listini", "Listini (Prezzi) - optional",
                                 self._handle_listini_upload)
        listini_row.grid(row=0, column=1, padx=2, pady=1, sticky="ew")
        self.listini_row = listini_row

    def _build_toolbar(self):
        bar = ttk.Frame(self)
        bar.pack(side="top", fill="x", padx=8, pady=(1, 2))

        self.refresh_btn = ttk.Button(bar, text="2) Refresh", command=self._on_refresh)
        self.refresh_btn.pack(side="left", padx=4)

        self.upload_data_btn = ttk.Button(bar, text="Upload Data", command=self._on_upload_data)
        self.upload_data_btn.pack(side="left", padx=4)

        self.export_btn = ttk.Button(bar, text="📤 Export to Excel", command=self._on_export)
        self.export_btn.pack(side="left", padx=4)

        self.abbina_btn = ttk.Button(bar, text="Da abbinare", command=self._open_abbina)
        self.abbina_btn.pack(side="left", padx=4)

        self.yarn_shortage_btn = ttk.Button(bar, text="Mancanza Filato", command=self._open_yarn_shortage)
        self.yarn_shortage_btn.pack(side="left", padx=4)

        self.copertura_btn = ttk.Button(bar, text="Copertura", command=self._open_copertura)
        self.copertura_btn.pack(side="left", padx=4)

        self.on_time_btn = ttk.Button(bar, text="On-Time Delivery", command=self._open_on_time_delivery)
        self.on_time_btn.pack(side="left", padx=4)

        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", self._on_search_changed)
        ttk.Label(bar, text="Search:").pack(side="left", padx=(14, 4))
        self.search_entry = ttk.Entry(bar, textvariable=self.search_var, width=28)
        self.search_entry.pack(side="left")
        self.clear_btn = ttk.Button(bar, text="Clear", command=self._on_clear_search)
        self.clear_btn.pack(side="left", padx=(4, 6))

        self.summary_lbl = ttk.Label(bar, text="")
        self.summary_lbl.pack(side="right", padx=8)

    def _open_copertura(self):
        """Show and export the dyeing coverage summary for machines 3–12."""
        if self._focus_child_window("copertura"):
            return
        copertura = self.loaded_frames.get("copertura")
        if not isinstance(copertura, pd.DataFrame) or copertura.empty:
            messagebox.showinfo("Copertura", "Upload the Copertura file first.", parent=self)
            return
        if "machine" not in copertura.columns or self.current_df.empty:
            messagebox.showinfo("Copertura", "Refresh Situazione data first.", parent=self)
            return

        situation = self.current_df.copy()

        for frame in (situation, copertura):
            frame["bagno"] = frame["bagno"].fillna("").astype(str).str.strip()
            frame["bagno_key"] = frame["bagno"].map(business_logic.bagno_key)
        merged = situation.merge(
            copertura[["bagno_key", "machine"]].drop_duplicates("bagno_key"),
            on="bagno_key", how="inner",
        )

        merged["machine_number"] = merged["machine"].map(business_logic.machine_number_from_label)
        merged = merged[merged["machine_number"].between(3, 12, inclusive="both")]
        if merged.empty:
            messagebox.showinfo("Copertura", "No Situazione colours match machines 3–12.", parent=self)
            return

        coverage_until = business_logic.machine_coverage_until

        def build_summary(view):
            columns = ["machine_number", "cliente", "total_colors", "pgx", "available", "covered_until"]
            if view.empty:
                return pd.DataFrame(columns=columns)
            work = view.copy()
            work["is_pgx"] = work["comment"].fillna("").astype(str).str.upper().str.startswith("PG-X")
            summary = work.groupby(["machine_number", "cliente"], dropna=False).agg(
                total_colors=("colore", "size"), pgx=("is_pgx", "sum"),
            ).reset_index()
            summary["available"] = summary["total_colors"] - summary["pgx"]
            totals = summary.groupby("machine_number")["total_colors"].sum().to_dict()
            summary["covered_until"] = summary["machine_number"].map(lambda x: coverage_until(totals.get(x, 0)))
            return summary

        window = tk.Toplevel(self)
        keep_window_on_top(window)
        self._child_windows["copertura"] = window
        window.title("Copertura — macchine 3–12")
        window.geometry("1050x600")
        window.minsize(800, 450)
        # Keep the normal Windows title bar so the native minimize, maximize,
        # restore, and close buttons remain available beside X.
        window.resizable(True, True)
        ttk.Label(window, text="Copertura: 2 colori al giorno per macchina — venerdì escluso", font=("Segoe UI", 10, "bold")).pack(anchor="w", padx=10, pady=(10, 4))
        search_row = ttk.Frame(window)
        search_row.pack(fill="x", padx=10, pady=(0, 6))
        ttk.Label(search_row, text="Filtra per:").pack(side="left", padx=(0, 6))
        filter_kind = tk.StringVar(value="Tutti")
        filter_value = tk.StringVar(value="Tutti")
        kind_combo = ttk.Combobox(
            search_row, textvariable=filter_kind, state="readonly", width=14,
            values=("Tutti", "Macchina", "Cliente"),
        )
        kind_combo.pack(side="left", padx=(0, 6))
        value_combo = ttk.Combobox(search_row, textvariable=filter_value, state="readonly", width=24)
        value_combo.pack(side="left")
        status = ttk.Label(search_row, text="")
        status.pack(side="right")

        frame = ttk.Frame(window)
        frame.pack(fill="both", expand=True, padx=10, pady=(0, 8))
        columns = ("machine", "cliente", "colors", "pgx", "available", "covered_until")
        labels = {"machine": "Macchina", "cliente": "Cliente", "colors": "Totale colori", "pgx": "Manca Filato (PG-X)", "available": "Disponibile", "covered_until": "Coperta fino al"}
        tree = ttk.Treeview(frame, columns=columns, show="headings")
        widths = [85, 130, 110, 140, 105, 150]
        for column, width in zip(columns, widths):
            tree.heading(column, text=labels[column])
            tree.column(column, width=width, anchor="center")
        yscroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=yscroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        tree.tag_configure("total", background="#e8f1fb", font=("Segoe UI", 9, "bold"))

        def current_view():
            selected_kind = filter_kind.get()
            selected_value = filter_value.get()
            if selected_kind == "Tutti" or selected_value == "Tutti":
                return merged
            if selected_kind == "Macchina":
                selected_machine = business_logic.machine_number_from_label(selected_value)
                return merged.loc[merged["machine_number"].eq(selected_machine)]
            return merged.loc[merged["cliente"].astype(str).eq(selected_value)]

        def refresh_filter_values(*_args):
            selected_kind = filter_kind.get()
            if selected_kind == "Macchina":
                values = ["Tutti"] + [
                    business_logic.machine_label(int(number))
                    for number in sorted(merged["machine_number"].dropna().unique())
                ]
            elif selected_kind == "Cliente":
                values = ["Tutti"] + sorted(merged["cliente"].dropna().astype(str).unique())
            else:
                values = ["Tutti"]
            value_combo.configure(values=values)
            filter_value.set("Tutti")
            render()

        def render():
            view = current_view()
            summary = build_summary(view)
            tree.delete(*tree.get_children())
            machines = range(3, 13)
            if filter_kind.get() == "Macchina" and filter_value.get() != "Tutti":
                selected_machine = business_logic.machine_number_from_label(filter_value.get())
                machines = [selected_machine] if selected_machine is not None else []
            for machine in machines:
                rows = summary[summary["machine_number"] == machine]
                total = int(rows["total_colors"].sum()) if not rows.empty else 0
                until = coverage_until(total)
                machine_label = business_logic.machine_label(machine)
                if rows.empty:
                    tree.insert("", "end", values=(machine_label, "-", 0, 0, 0, "-"))
                else:
                    for _, row in rows.sort_values("cliente").iterrows():
                        tree.insert("", "end", values=(machine_label, row["cliente"], int(row["total_colors"]), int(row["pgx"]), int(row["available"]), row["covered_until"]))
                    tree.insert("", "end", values=(machine_label, "TOTALE", total, int(rows["pgx"].sum()), total - int(rows["pgx"].sum()), until), tags=("total",))
            status.config(text=f"{len(view):,} colori")

        def export_summary():
            path = filedialog.asksaveasfilename(parent=window, title="Esporta Copertura", defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")], initialfile="copertura.xlsx")
            if not path:
                return
            view = current_view()
            summary = build_summary(view)
            rows = []
            machines = range(3, 13)
            if filter_kind.get() == "Macchina" and filter_value.get() != "Tutti":
                selected_machine = business_logic.machine_number_from_label(filter_value.get())
                machines = [selected_machine] if selected_machine is not None else []
            for machine in machines:
                part = summary[summary["machine_number"] == machine]
                total = int(part["total_colors"].sum()) if not part.empty else 0
                machine_label = business_logic.machine_label(machine)
                if part.empty:
                        rows.append([machine_label, "-", 0, 0, 0, "-"])
                else:
                    for _, row in part.sort_values("cliente").iterrows():
                        rows.append([machine_label, row["cliente"], int(row["total_colors"]), int(row["pgx"]), int(row["available"]), row["covered_until"]])
                    pgx_total = int(part["pgx"].sum())
                    rows.append([machine_label, "TOTALE", total, pgx_total, total - pgx_total, coverage_until(total)])
            try:
                pd.DataFrame(rows, columns=[labels[c] for c in columns]).to_excel(path, index=False, sheet_name="Copertura")
                messagebox.showinfo("Copertura", f"Export completato:\n{path}", parent=window)
            except Exception as exc:  # noqa: BLE001
                messagebox.showerror("Copertura", str(exc), parent=window)

        kind_combo.bind("<<ComboboxSelected>>", refresh_filter_values)
        value_combo.bind("<<ComboboxSelected>>", lambda _event: render())
        buttons = ttk.Frame(window)
        buttons.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Button(buttons, text="📤 Export to Excel", command=export_summary).pack(side="right", padx=3)
        refresh_filter_values()
        render()

    def _open_on_time_delivery(self):
        """Per-client on-time delivery score: Consegna (promised) vs Data
        Uscita (actual), across every Partita persisted in the local DB --
        not just what's currently loaded, so this reflects the full history
        of shipped batches, not only today's snapshot."""
        if self._focus_child_window("on_time"):
            return
        states = db.get_all_states()
        summary = reports.compute_on_time_delivery(states)
        if summary.empty:
            messagebox.showinfo(
                "On-Time Delivery",
                "No shipped batches with both a Consegna and a Data Uscita date yet.",
                parent=self,
            )
            return

        window = tk.Toplevel(self)
        keep_window_on_top(window)
        self._child_windows["on_time"] = window
        window.title("On-Time Delivery by Client")
        window.geometry("780x480")
        window.minsize(600, 350)
        ttk.Label(
            window, text="Share of shipped batches that left on or before the promised Consegna date",
            font=("Segoe UI", 10, "bold"),
        ).pack(anchor="w", padx=10, pady=(10, 6))

        frame = ttk.Frame(window)
        frame.pack(fill="both", expand=True, padx=10, pady=(0, 8))
        columns = ("cliente", "shipped", "on_time", "late", "on_time_pct", "avg_delay_days")
        labels = {
            "cliente": "Cliente", "shipped": "Shipped", "on_time": "On Time",
            "late": "Late", "on_time_pct": "On-Time %", "avg_delay_days": "Avg Delay (days, when late)",
        }
        tree = ttk.Treeview(frame, columns=columns, show="headings")
        widths = [160, 80, 80, 70, 90, 190]
        for column, width in zip(columns, widths):
            tree.heading(column, text=labels[column])
            tree.column(column, width=width, anchor="center")
        yscroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=yscroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        tree.tag_configure("bad", foreground="#b91c1c")
        tree.tag_configure("good", foreground="#15803d")

        for _, row in summary.iterrows():
            tag = "bad" if row["on_time_pct"] < 70 else ("good" if row["on_time_pct"] >= 95 else "")
            tree.insert("", "end", values=(
                row["cliente"], int(row["shipped"]), int(row["on_time"]), int(row["late"]),
                f"{row['on_time_pct']:.1f}%", f"{row['avg_delay_days']:.1f}",
            ), tags=(tag,) if tag else ())

        def export_summary():
            path = filedialog.asksaveasfilename(
                parent=window, title="Export On-Time Delivery", defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")], initialfile="on_time_delivery.xlsx",
            )
            if not path:
                return
            try:
                summary.rename(columns=labels).to_excel(path, index=False, sheet_name="On-Time Delivery")
                messagebox.showinfo("On-Time Delivery", f"Export completed:\n{path}", parent=window)
            except Exception as exc:  # noqa: BLE001
                messagebox.showerror("On-Time Delivery", str(exc), parent=window)

        buttons = ttk.Frame(window)
        buttons.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Button(buttons, text="📤 Export to Excel", command=export_summary).pack(side="right", padx=3)

    def _focus_child_window(self, key):
        """Focus an already-open child window instead of opening a duplicate."""
        window = self._child_windows.get(key)
        if window is None:
            return False
        try:
            if not window.winfo_exists():
                self._child_windows.pop(key, None)
                return False
            if window.state() == "iconic":
                window.deiconify()
            window.lift()
            window.focus_force()
            return True
        except tk.TclError:
            self._child_windows.pop(key, None)
            return False

    def _build_treeview(self):
        cols = [key for key, _, _ in COLUMN_SPEC]
        self.columns = cols

        frame = ttk.Frame(self, borderwidth=1, relief="solid")
        frame.pack(side="top", fill="both", expand=True, padx=8, pady=(0, 6))

        self.tree = ttk.Treeview(frame, columns=cols, show="headings", selectmode="browse",
                                 style="Situazione.Treeview")
        for c in cols:
            ctype = COLUMN_TYPES.get(c, "text")
            anchor = "center" if ctype in ("number", "date") else "w"
            self.tree.heading(c, text=COLUMN_LABELS.get(c, c), anchor="center",
                              command=lambda c=c: self._sort_by(c))
            self.tree.column(c, width=100, minwidth=60, anchor=anchor, stretch=False)

        vsb = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)

        for status, color in STATUS_COLORS.items():
            self.tree.tag_configure(status, background=color)
        self.tree.tag_configure("Ritinta", background="#d9c6f0")
        self.tree.tag_configure("price_error", background="#fee2e2", foreground="#991b1b")
        self.tree.tag_configure("stripe", background="#f3f6fa")
        self.tree.tag_configure("normal", background="#ffffff")

    def _autosize_columns(self, df):
        """Fit columns to visible content while keeping the table usable."""
        body_font = tkfont.Font(family="Segoe UI", size=9)
        heading_font = tkfont.Font(family="Segoe UI", size=9, weight="bold")
        for col in self.columns:
            values = []
            if not df.empty and col in df.columns:
                values = [str(value) for value in df[col].fillna("").tolist()]
            # Measuring only the longest string is much faster than measuring
            # every cell, especially for large Excel exports.
            longest_value = max(values, key=len, default="")
            longest = max(heading_font.measure(COLUMN_LABELS.get(col, col)), body_font.measure(longest_value))
            width = max(65, min(longest + 22, 300))
            self.tree.column(col, width=width, minwidth=min(width, 65), stretch=False)

    # ------------------------------------------------------------- shared DFM / uploads














    # -------------------------------------------------------------- refresh













    # -------------------------------------------------------------- display
    def _render_tree(self, df, autosize=True):
        """Render rows in small UI batches so large snapshots do not freeze Tk."""
        self._tree_render_generation += 1
        generation = self._tree_render_generation
        if self._tree_render_after_id is not None:
            try:
                self.after_cancel(self._tree_render_after_id)
            except tk.TclError:
                pass
            self._tree_render_after_id = None

        self.tree.delete(*self.tree.get_children())
        if autosize:
            self._autosize_columns(df)
        if df.empty:
            return
        display_df = df.reindex(columns=self.columns, fill_value="")
        for date_column in ("data", "delivery_date", "consegna", "tinto", "data_qualita", "data_uscita"):
            if date_column in display_df.columns:
                original = display_df[date_column].fillna("").astype(str)
                parsed = pd.to_datetime(display_df[date_column], errors="coerce")
                formatted = parsed.dt.strftime("%d/%m/%Y")
                display_df[date_column] = formatted.where(parsed.notna(), original)
        rows = list(display_df.itertuples(index=False, name=None))
        prezzo_idx = self.columns.index("prezzo") if "prezzo" in self.columns else None
        prezzo_lisini_idx = self.columns.index("prezzo_lisini") if "prezzo_lisini" in self.columns else None

        def _is_price_error(row_vals):
            if prezzo_idx is None or prezzo_lisini_idx is None:
                return False
            p_val, pl_val = row_vals[prezzo_idx], row_vals[prezzo_lisini_idx]
            try:
                p_clean = str(p_val or "").replace("$", "").replace("USD", "").replace("usd", "").strip().replace(",", ".")
                p_num = float(p_clean) if p_clean else None
            except (TypeError, ValueError):
                p_num = None
            try:
                pl_clean = str(pl_val or "").replace("$", "").replace("USD", "").replace("usd", "").strip().replace(",", ".")
                pl_num = float(pl_clean) if pl_clean else None
            except (TypeError, ValueError):
                pl_num = None

            if p_num is not None and abs(p_num - 0.01) < 1e-9:
                return True
            if p_num is not None and pl_num is not None and abs(p_num - pl_num) > 0.01:
                return True
            return False

        def insert_chunk(start=0):
            if generation != self._tree_render_generation:
                return
            end = min(start + 150, len(rows))
            for index in range(start, end):
                values = rows[index]
                status = values[self.columns.index("new_comment")]
                tag = "Ritinta" if str(status).startswith("Ritinta") else status
                if _is_price_error(values):
                    tag = "price_error"
                elif not tag:
                    tag = "stripe" if index % 2 else ""
                self.tree.insert("", "end", values=values, tags=((tag,) if tag else ()))
            if end < len(rows):
                self._tree_render_after_id = self.after(1, insert_chunk, end)
            else:
                self._tree_render_after_id = None

        insert_chunk()

    def _on_search_changed(self, *_args):
        """Debounce typing so the whole table is not redrawn per keystroke."""
        if self._filter_after_id is not None:
            self.after_cancel(self._filter_after_id)
        self._filter_after_id = self.after(180, self._apply_filter)

    def _apply_filter(self):
        self._filter_after_id = None
        q = self.search_var.get().strip().lower()
        if not q:
            self._render_tree(self.current_df)
            return
        if self.current_df.empty:
            return
        searchable = self.current_df.reindex(columns=self.columns, fill_value="").astype(str)
        mask = searchable.apply(lambda col: col.str.contains(q, case=False, regex=False)).any(axis=1)
        self._render_tree(self.current_df[mask], autosize=False)

    def _on_clear_search(self):
        """Clear search input and restore default treeview display/sorting."""
        if self._filter_after_id is not None:
            self.after_cancel(self._filter_after_id)
            self._filter_after_id = None
        self.search_var.set("")
        if not self.current_df.empty and "bagno" in self.current_df.columns:
            self.sort_state.clear()
            self.current_df = self.current_df.sort_values(
                by="bagno", ascending=True, key=lambda s: s.astype(str)
            )
            self.sort_state["bagno"] = False
        self._render_tree(self.current_df)

    def _sort_by(self, col):
        ascending = self.sort_state.get(col, True)
        if self.current_df.empty:
            return
        self.current_df = self.current_df.sort_values(by=col, ascending=ascending, key=lambda s: s.astype(str))
        self.sort_state[col] = not ascending
        self._apply_filter()

    def _open_abbina(self):
        if self._focus_child_window("abbina"):
            return
        suggestions = build_suggestions(self.current_df, max_extra_percent=0.20)
        window = tk.Toplevel(self)
        keep_window_on_top(window)
        self._child_windows["abbina"] = window
        window.title("Da abbinare")
        window.geometry("1250x600")
        window.minsize(850, 350)
        window.resizable(True, True)

        top = ttk.Frame(window)
        top.pack(fill="x", padx=8, pady=8)
        ttk.Label(top, text=(f"{len(suggestions)} righe — stesso Codice/Colore, Titolo compatibile; limite extra 20% (oltre = ⚠)"),
                  foreground="#344054").pack(side="left")
        ttk.Label(top, text="Search:").pack(side="left", padx=(18, 4))
        search_var = tk.StringVar()
        search_entry = ttk.Entry(top, textvariable=search_var, width=24)
        search_entry.pack(side="left")
        ttk.Button(top, text="Clear", width=6, command=lambda: search_var.set("")).pack(side="left", padx=(4, 0))
        ttk.Button(top, text="Export Abbina", command=lambda: self._export_abbina(suggestions)).pack(side="right")

        cols = ["titolo", "codice", "colore", "rocche", "partita", "bagno", "abbina",
                "tot_rocche", "mc_target", "polmoni", "extra_percent", "motivo", "new_comment"]
        labels = {"titolo": "Titolo", "codice": "Codice", "colore": "Colore", "rocche": "Rocche",
                  "partita": "Partita", "bagno": "Bagno", "abbina": "Abbina", "tot_rocche": "Tot. Rocche",
                  "mc_target": "Capacità", "polmoni": "Polmoni", "extra_percent": "Extra %",
                  "motivo": "Motivo", "new_comment": "New Comment"}
        frame = ttk.Frame(window)
        frame.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        style = ttk.Style(window)
        style.configure("Abbina.Treeview", rowheight=28, font=("Segoe UI", 9))
        tree = ttk.Treeview(frame, columns=cols, show="headings", style="Abbina.Treeview")
        for col in cols:
            tree.heading(col, text=labels[col])
            tree.column(col, width=105 if col not in ("motivo", "abbina", "new_comment") else 230, anchor="center")
        yscroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        xscroll = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        tree.tag_configure("group_a", background="#eaf2f8")
        tree.tag_configure("group_b", background="#fff7e6")

        def render():
            tree.delete(*tree.get_children())
            query = search_var.get().strip().casefold()
            visible = suggestions
            if query:
                mask = suggestions[cols].fillna("").astype(str).apply(
                    lambda column: column.str.casefold().str.contains(query, regex=False)
                ).any(axis=1)
                visible = suggestions[mask]
            group_tags = {
                group: "group_a" if index % 2 == 0 else "group_b"
                for index, group in enumerate(
                    suggestions.apply(
                        lambda row: f"{row.get('codice', '')}|{row.get('colore', '')}|{row.get('motivo', '')}",
                        axis=1,
                    ).drop_duplicates()
                )
            }
            for _, row in visible.iterrows():
                group = f"{row.get('codice', '')}|{row.get('colore', '')}|{row.get('motivo', '')}"
                values = [f"{row[c]:.1%}" if c == "extra_percent" else row.get(c, "") for c in cols]
                tree.insert("", "end", values=values, tags=(group_tags[group],))

        search_var.trace_add("write", lambda *_: render())
        render()

    def _export_abbina(self, suggestions):
        if suggestions.empty:
            messagebox.showinfo("Da abbinare", "Non ci sono combinazioni entro il limite del 20%.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                             initialfile="Da_abbinare.xlsx")
        if not path:
            return
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        headers = ["Titolo", "Codice", "Colore", "Rocche", "Partita", "Bagno", "Abbina",
                   "Tot. Rocche", "Capacità M/C", "Polmoni", "Extra %", "Motivo", "New Comment"]
        wb = Workbook()
        ws = wb.active
        ws.title = "Da abbinare"
        ws.append(headers)

        header_fill = PatternFill("solid", fgColor="16324F")
        header_font = Font(name="Arial", bold=True, color="FFFFFF")
        thin = Side(style="thin", color="B8C6D6")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border

        group_colors = {"group_a": "EAF2F8", "group_b": "FFF7E6"}
        group_tags = {
            group: "group_a" if index % 2 == 0 else "group_b"
            for index, group in enumerate(
                suggestions.apply(
                    lambda row: f"{row.get('codice', '')}|{row.get('colore', '')}|{row.get('motivo', '')}",
                    axis=1,
                ).drop_duplicates()
            )
        }
        for _, row in suggestions.iterrows():
            group = f"{row.get('codice', '')}|{row.get('colore', '')}|{row.get('motivo', '')}"
            values = [row.get("titolo", ""), row.get("codice", ""), row.get("colore", ""),
                      row.get("rocche", ""), row.get("partita", ""), row.get("bagno", ""),
                      row.get("abbina", ""), row.get("tot_rocche", ""), row.get("mc_target", ""),
                      row.get("polmoni", ""), f"{row.get('extra_percent', 0):.1%}",
                      row.get("motivo", ""), row.get("new_comment", "")]
            ws.append(values)
            fill = PatternFill("solid", fgColor=group_colors[group_tags[group]])
            for cell in ws[ws.max_row]:
                cell.fill = fill
                cell.font = Font(name="Arial")
                cell.alignment = center
                cell.border = border

        last_row = ws.max_row
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"
        ws.freeze_panes = "A2"
        for column_index, header in enumerate(headers, start=1):
            letter = get_column_letter(column_index)
            values = [str(ws.cell(row=row, column=column_index).value or "") for row in range(1, last_row + 1)]
            width = min(max(max(len(value) for value in values) + 3, len(header) + 2, 10), 45)
            ws.column_dimensions[letter].width = width
        ws.row_dimensions[1].height = 28
        safe_save_workbook(wb, path)
        messagebox.showinfo("Completed", f"Export completed successfully:\n{path}")

    def _open_yarn_shortage(self):
        """Opens Mancanza Filato as a popup window (same pattern as Da abbinare)."""
        if self._focus_child_window("yarn_shortage"):
            return
        window = tk.Toplevel(self)
        keep_window_on_top(window)
        self._child_windows["yarn_shortage"] = window
        window.title("Mancanza Filato")
        window.geometry("1100x650")
        window.minsize(750, 350)
        window.resizable(True, True)

        shortage_view = YarnShortageTab(window, self)
        shortage_view.pack(fill="both", expand=True)
        self.add_table_loaded_callback(shortage_view.refresh)

        def on_close():
            try:
                self._table_loaded_callbacks.remove(shortage_view.refresh)
            except ValueError:
                pass
            self._child_windows.pop("yarn_shortage", None)
            window.destroy()

        window.protocol("WM_DELETE_WINDOW", on_close)
        shortage_view.refresh()

    # -------------------------------------------------------------- export
    @staticmethod
    def _to_number(text):
        """Parse a stored value into a real number. Handles European comma
        decimals (e.g. '128,00') as well as plain '128'. Returns None for
        blank/unparseable values so the cell stays genuinely empty."""
        if text is None:
            return None
        s = str(text).strip()
        if not s or s.lower() == "nan":
            return None
        s = s.replace(".", "").replace(",", ".") if ("," in s and s.count(",") == 1) else s
        try:
            value = float(s)
        except ValueError:
            return None
        return int(value) if value.is_integer() else value

    @staticmethod
    def _to_date(text):
        """Parse a stored 'YYYY-MM-DD' string into a real date. Returns None
        for blank/unparseable values."""
        if not text:
            return None
        s = str(text).strip()
        if not s or s.lower() == "nan":
            return None
        try:
            return datetime.strptime(s, "%Y-%m-%d").date()
        except ValueError:
            return None

    def _on_export(self):
        if self.current_df.empty:
            messagebox.showinfo("No data", "Refresh the data before exporting.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                             filetypes=[("Excel", "*.xlsx")],
                                             initialfile="Situazione_Generale.xlsx")
        if not path:
            return

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import CellIsRule, FormulaRule

        # Export exactly the rows currently visible after Search.  This keeps
        # Excel consistent with the user's filtered Treeview.
        export_df = self.current_df.copy()
        q = self.search_var.get().strip().lower()
        if q and not export_df.empty:
            searchable = export_df.reindex(columns=self.columns, fill_value="").astype(str)
            mask = searchable.apply(lambda col: col.str.contains(q, case=False, regex=False)).any(axis=1)
            export_df = export_df[mask]

        thin = Side(style="thin", color="B0B0B0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center")
        header_fill = PatternFill(start_color="16324F", end_color="16324F", fill_type="solid")

        wb = Workbook()
        ws = wb.active
        ws.title = "Situazione"
        headers = [header for _, header, _ in COLUMN_SPEC]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, name="Arial", color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        date_format = "dd/mm/yyyy"

        for _, r in export_df.iterrows():
            row_values = []
            for key, _header, ctype in COLUMN_SPEC:
                raw = r.get(key, "")
                if ctype == "number":
                    row_values.append(self._to_number(raw))
                elif ctype == "date":
                    row_values.append(self._to_date(raw))
                else:
                    row_values.append(raw if raw not in (None, "") else None)
            ws.append(row_values)

            status = str(r.get("new_comment", ""))
            color = color_for_status(status if not status.startswith("Ritinta") else "Ritinta")
            fill = PatternFill(start_color=color.replace("#", ""), end_color=color.replace("#", ""), fill_type="solid")
            for col_index, (key, _header, ctype) in enumerate(COLUMN_SPEC, start=1):
                cell = ws.cell(row=ws.max_row, column=col_index)
                cell.fill = fill
                cell.font = Font(name="Arial")
                cell.alignment = center
                cell.border = border
                if ctype == "date" and cell.value is not None:
                    cell.number_format = date_format

        last_row = ws.max_row
        last_col_letter = get_column_letter(len(COLUMN_SPEC))
        ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"
        ws.freeze_panes = "A2"
        col_letter = {key: get_column_letter(i) for i, (key, _h, _t) in enumerate(COLUMN_SPEC, start=1)}

        if last_row > 1:
            # Custom == "Check" -> red
            rng = f"{col_letter['custom']}2:{col_letter['custom']}{last_row}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Check"'], fill=red_fill))

            # Bagno duplicated anywhere in the column -> red
            rng = f"{col_letter['bagno']}2:{col_letter['bagno']}{last_row}"
            first = f"{col_letter['bagno']}2"
            formula = f'AND({first}<>"",COUNTIF(${col_letter["bagno"]}$2:${col_letter["bagno"]}${last_row},{first})>1)'
            ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=red_fill))

            # Days in Q.C > 5 -> red
            rng = f"{col_letter['days_in_qc']}2:{col_letter['days_in_qc']}{last_row}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["5"], fill=red_fill))

            # Consegna before today (overdue) -> red
            rng = f"{col_letter['consegna']}2:{col_letter['consegna']}{last_row}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["TODAY()"], fill=red_fill))

            # Densita`(360-390) outside its expected range -> red
            rng = f"{col_letter['densita']}2:{col_letter['densita']}{last_row}"
            first = f"{col_letter['densita']}2"
            formula = f'AND({first}<>"",OR({first}<360,{first}>390))'
            ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=red_fill))

            # Ritardo (gg) > 0 (delivery delay) -> red
            rng = f"{col_letter['ritardo_consegna']}2:{col_letter['ritardo_consegna']}{last_row}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=red_fill))

        for key, header, _ctype in COLUMN_SPEC:
            letter = col_letter[key]
            values = [str(v) for v in export_df[key].tolist()] if key in export_df.columns else []
            longest = max([len(header)] + [len(v) for v in values]) if values else len(header)
            ws.column_dimensions[letter].width = min(max(longest + 2, 10), 35)

        safe_save_workbook(wb, path)
        logger.info("Situazione: exported %d visible rows to %s", len(export_df), path)
        messagebox.showinfo("Completed", f"Export completed successfully:\n{path}")


if __name__ == "__main__":
    root = tk.Tk()
    root.title("Situazione (standalone test)")
    root.geometry("1400x760")
    tab = SituazioneTab(root)
    tab.pack(fill="both", expand=True)
    root.mainloop()
