from __future__ import annotations

import re
import threading
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from utility.master_data import raw_articolo_for
from utility.path_manager import source_path
from utility.utils import keep_window_on_top, bind_escape_to_close
from .biglietti_exports import (
    RawYarnMatch,
    _filato_rows,
    delete_pg_x_row,
    delete_shipped_shared_rows,
    export_filato_full,
    export_word,
    load_create_excel_records,
    move_pg_x_to_orders,
    pg_x_batch_stock_status,
    remove_uploaded_pgx_rows,
    save_pg_x_partita,
    update_order_row,
    update_pg_x_row,
)


def _upload_partita_key(value):
    text = " ".join(str(value or "").split())
    try:
        number = float(text.replace(",", "."))
    except ValueError:
        return text.casefold()
    return str(int(number)) if number.is_integer() else str(number)


class SharedOrdersActionsMixin:
    @staticmethod
    def _toggle_child_maximize(window: tk.Toplevel) -> None:
        try:
            window.state("normal" if window.state() == "zoomed" else "zoomed")
        except tk.TclError:
            pass
    @staticmethod
    def _start_child_drag(window: tk.Toplevel, event) -> None:
        window._drag_origin = (event.x_root, event.y_root, window.winfo_x(), window.winfo_y())
    @staticmethod
    def _drag_child(window: tk.Toplevel, event) -> None:
        origin = getattr(window, "_drag_origin", None)
        if origin is None:
            return
        start_x, start_y, window_x, window_y = origin
        window.geometry(f"+{window_x + event.x_root - start_x}+{window_y + event.y_root - start_y}")
    @staticmethod
    def _restore_child_window(window: tk.Toplevel) -> None:
        try:
            window.state("normal")
            window.deiconify()
        except tk.TclError:
            pass
    def _save_pg_x_details_in_background(
        self, window, partita_col, values, datasets, refresh_records, rebuild,
        partita_gg_var, sheet_name="PG-X", allow_article_mismatch=False,
    ):
        """Persist the PG-X editor fields, then apply the normal Partita GG flow."""
        if getattr(self, "_pg_x_saving", False):
            return
        self._pg_x_saving = True

        def as_number(text):
            value = str(text or "").strip()
            if not value:
                return ""
            try:
                number = float(value.replace(",", "."))
                return int(number) if number.is_integer() else number
            except ValueError:
                return value

        updates = {
            "Articolo": values.get("Articolo", ""),
            "Titolo": values.get("Titolo", ""),
            "Print": values.get("Print", ""),
            "Partita GG": values.get("Partita GG", ""),
            "Bagno": values.get("Bagno", ""),
            "Rocche": as_number(values.get("Rocche", "")),
            "M/C": values.get("M/C", ""),
        }

        def worker():
            try:
                partita_gg = str(values.get("Partita GG", "")).strip()
                if sheet_name == "Orders":
                    magazino_summary = None
                    if partita_gg:
                        _codes, _densita_map, _vmm_ratio_map, _prices, magazino_summary = self._load_common_sources()
                    result = update_order_row(
                        self.shared_excel_path, partita_col, updates,
                        magazino_summary=magazino_summary,
                        allow_article_mismatch=allow_article_mismatch,
                    )
                    if result["moved_to_pgx"]:
                        message = f"Moved {result['updated']} row(s) to PG-X."
                    else:
                        message = f"Updated {result['updated']} row(s) in Orders."
                else:
                    update_pg_x_row(self.shared_excel_path, partita_col, updates)
                    if partita_gg:
                        _codes, densita_map, vmm_ratio_map, _prices, magazino_summary = self._load_common_sources()
                        result = save_pg_x_partita(
                            self.shared_excel_path, partita_col, partita_gg,
                            densita_map=densita_map, vmm_ratio_map=vmm_ratio_map,
                            magazino_summary=magazino_summary,
                            allow_article_mismatch=allow_article_mismatch,
                        )
                        message = f"Moved {result['updated']} row(s) to Orders. Available raw yarn: {result['available']:g} rocche."
                    else:
                        result = {"updated": 1}
                        message = "PG-X color updated."
                new_datasets = {
                    "Orders": load_create_excel_records(self.shared_excel_path, sheet_name="Orders"),
                    "PG-X": load_create_excel_records(self.shared_excel_path, sheet_name="PG-X"),
                }
                self.after(0, lambda: finish(new_datasets, message))
            except Exception as exc:
                self.after(0, lambda exc=exc: fail(exc))

        def finish(new_datasets, message):
            self._pg_x_saving = False
            refresh_records(new_datasets)
            partita_gg_var.set("")
            rebuild()
            messagebox.showinfo("PG-X Saved", message, parent=window)

        def fail(exc):
            self._pg_x_saving = False
            text = str(exc)
            if self._on_notification and ("does not belong to article" in text or "open or locked" in text or "Cannot save" in text):
                self._on_notification("pgx-save-error", "PG-X save needs attention", text, "Create (EXCEL+Biglietti)", "high")
            messagebox.showerror("PG-X Save Error", str(exc), parent=window)

        threading.Thread(target=worker, daemon=True).start()
    def _save_pg_x_in_background(self, window, partita_col, partita_gg, datasets, refresh_records, rebuild, partita_gg_var, allow_article_mismatch=False):
        if getattr(self, "_pg_x_saving", False):
            return
        self._pg_x_saving = True

        def worker():
            try:
                _codes, densita_map, vmm_ratio_map, _prices, magazino_summary = self._load_common_sources()
                result = save_pg_x_partita(
                    self.shared_excel_path, partita_col, partita_gg,
                    densita_map=densita_map, vmm_ratio_map=vmm_ratio_map,
                    magazino_summary=magazino_summary,
                    allow_article_mismatch=allow_article_mismatch,
                )
                new_datasets = {
                    "Orders": load_create_excel_records(self.shared_excel_path, sheet_name="Orders"),
                    "PG-X": load_create_excel_records(self.shared_excel_path, sheet_name="PG-X"),
                }
                self.after(0, lambda: finish(result, new_datasets))
            except Exception as exc:
                self.after(0, lambda exc=exc: fail(exc))

        def finish(result, new_datasets):
            self._pg_x_saving = False
            refresh_records(new_datasets)
            partita_gg_var.set("")
            rebuild()
            messagebox.showinfo(
                "PG-X Saved",
                f"Updated {result['updated']} row(s). Available raw yarn: {result['available']:g} rocche.",
                parent=window,
            )

        def fail(exc):
            self._pg_x_saving = False
            text = str(exc)
            if self._on_notification and ("does not belong to article" in text or "open or locked" in text or "Cannot save" in text):
                self._on_notification("pgx-save-error", "PG-X save needs attention", text, "Create (EXCEL+Biglietti)", "high")
            messagebox.showerror("PG-X Save Error", str(exc), parent=window)

        threading.Thread(target=worker, daemon=True).start()
    def _pg_x_row_action_in_background(self, window, action, partita_col, datasets, refresh_records, rebuild):
        if getattr(self, "_pg_x_action_running", False):
            return
        self._pg_x_action_running = True

        def worker():
            try:
                count = move_pg_x_to_orders(self.shared_excel_path, partita_col) if action == "move" else delete_pg_x_row(self.shared_excel_path, partita_col)
                new_datasets = {
                    "Orders": load_create_excel_records(self.shared_excel_path, sheet_name="Orders"),
                    "PG-X": load_create_excel_records(self.shared_excel_path, sheet_name="PG-X"),
                }
                self.after(0, lambda: finish(count, new_datasets))
            except Exception as exc:
                self.after(0, lambda exc=exc: fail(exc))

        def finish(count, new_datasets):
            self._pg_x_action_running = False
            refresh_records(new_datasets)
            rebuild()
            messagebox.showinfo(
                "PG-X",
                f"{count} row(s) {'sent to Orders' if action == 'move' else 'deleted'}.",
                parent=window,
            )

        def fail(exc):
            self._pg_x_action_running = False
            messagebox.showerror("PG-X Error", str(exc), parent=window)

        threading.Thread(target=worker, daemon=True).start()
    def _delete_shipped_colors(
        self, window, sheet_name, record_by_iid, sheet_by_iid,
        datasets, refresh_records, rebuild,
    ):
        if getattr(self, "_delete_shipped_running", False):
            return
        uscita_path = source_path("uscita", existing_only=True)
        if not uscita_path:
            return messagebox.showwarning(
                "Uscita File Required",
                "Upload the Uscita file in Situazione first, then try again.",
                parent=window,
            )
        self._delete_shipped_running = True

        def normalize(value):
            text = str(value or "").strip()
            try:
                number = float(text.replace(",", "."))
                return str(int(number)) if number.is_integer() else str(number)
            except ValueError:
                return text.casefold()

        def worker():
            try:
                from parsers.situazione_loaders import load_uscita
                uscita_df, errors = load_uscita(str(uscita_path))
                if errors or uscita_df is None or uscita_df.empty:
                    detail = "; ".join(errors) if errors else "The Uscita file has no valid shipped rows."
                    raise ValueError(detail)
                shipped = {normalize(value) for value in uscita_df["partita"].tolist() if normalize(value)}
                page_records = [
                    record for iid, record in record_by_iid.items()
                    if sheet_by_iid.get(iid) == sheet_name
                ]
                count = sum(1 for record in page_records if normalize(record.colored_batch) in shipped)
                self.after(0, lambda: confirm_and_delete(count, shipped))
            except Exception as exc:
                self.after(0, lambda exc=exc: shipped_delete_failed(exc, window))

        def confirm_and_delete(count, shipped):
            if not count:
                self._delete_shipped_running = False
                return messagebox.showinfo(
                    "Delete Shipped Colors",
                    f"No shipped colors were found in {sheet_name}.",
                    parent=window,
                )
            if not messagebox.askyesno(
                "Delete Shipped Colors",
                f"{count} color(s) in {sheet_name} have already been shipped.\n\nDelete them now?",
                parent=window,
            ):
                self._delete_shipped_running = False
                return

            def delete_worker():
                try:
                    deleted = delete_shipped_shared_rows(self.shared_excel_path, shipped, sheet_name)
                    new_datasets = {
                        "Orders": load_create_excel_records(self.shared_excel_path, sheet_name="Orders"),
                        "PG-X": load_create_excel_records(self.shared_excel_path, sheet_name="PG-X"),
                    }
                    self.after(0, lambda: delete_finished(deleted, new_datasets))
                except Exception as exc:
                    self.after(0, lambda exc=exc: shipped_delete_failed(exc, window))

            threading.Thread(target=delete_worker, daemon=True).start()

        def delete_finished(deleted, new_datasets):
            self._delete_shipped_running = False
            refresh_records(new_datasets)
            rebuild()
            messagebox.showinfo(
                "Delete Shipped Colors",
                f"Deleted {deleted} shipped color(s) from {sheet_name}.",
                parent=window,
            )

        def shipped_delete_failed(exc, parent):
            self._delete_shipped_running = False
            messagebox.showerror("Delete Shipped Colors", str(exc), parent=parent)

        threading.Thread(target=worker, daemon=True).start()
    def _print_selected_shared(
        self, window, selected: dict[str, bool], record_by_iid: dict[str, object],
        move_assigned_pg_x: bool = False,
    ):
        if not self.template_path or not self.template_path.is_file():
            return messagebox.showwarning("Missing Template", "Select the Biglietti.docx template first.", parent=window)
        records = [record_by_iid[iid] for iid, is_selected in selected.items() if is_selected]
        if not records:
            return messagebox.showwarning("No Orders Selected", "Select at least one color to print.", parent=window)
        # Keep one stable output file: every print replaces the previous
        # document with only the records passed in this time. Shared by
        # both "Print Selected Biglietti" (checked rows) and "Print
        # Assigned PG-X" (assigned rows) -- named neither "..._Selected"
        # (looked like the wrong button had run when Print Assigned PG-X
        # produced it) nor "Biglietti.docx" (that name is the *template*
        # file selected above, usually kept in this same folder -- reusing
        # it here would silently overwrite the user's own template).
        destination = self.shared_excel_path.parent / "Biglietti_Stampati.docx"
        pg_x_partita_cols = []
        if move_assigned_pg_x:
            pg_x_partita_cols = list(dict.fromkeys(
                str(record.colored_batch)
                for record in records
                if str(record.raw_batch or "").strip().upper().replace(" ", "") not in {"", "X", "PG-X", "PGX"}
            ))
        window.destroy()
        self.convert_btn.config(state="disabled")
        self._set_status("Creating selected Biglietti from shared Excel in progress...")
        threading.Thread(
            target=self._worker_shared_biglietti,
            args=(records, destination, pg_x_partita_cols),
            daemon=True,
        ).start()
    def _worker_shared_biglietti(self, records, destination: Path, pg_x_partita_cols=None):
        lines = []
        had_error = False
        try:
            export_word(destination, self.template_path, records, stem="Selected Orders")
            self._set_status(f"Created {len(records)} selected Biglietti.")
            lines.append(f"Created:\n{destination}")
        except Exception as exc:
            self._logger.exception("Shared Excel Biglietti (Word) failed")
            had_error = True
            lines.append(f"Biglietti (.docx) FAILED:\n{exc}")

        # Same file/logic as the old standalone "Filato X Tinturia" button:
        # the raw yarn to prepare for whichever rows are being printed here,
        # restricted to the ones that actually have a raw batch resolved
        # (Partita GG) -- a row still waiting for one has nothing to list
        # yet. Kept in its own try/except: a problem here (e.g. the Filato
        # X Tinturia.xlsx file is currently open in Excel) must not read as
        # "nothing happened" when the Biglietti above already printed fine,
        # and must not block the move-to-Orders step below either.
        filato_records = [
            record for record in records
            if str(record.raw_batch or "").strip().upper().replace(" ", "") not in {"", "X", "PG-X", "PGX"}
        ]
        if filato_records:
            try:
                _codes, _density, _vmm, _prices, summary = self._load_common_sources()
                rows = _filato_rows(filato_records, [], summary)
                filato_path = self.shared_excel_path.parent / "Filato X Tinturia.xlsx"
                export_filato_full(filato_path, [
                    RawYarnMatch(
                        articolo=str(row.get("Articolo", "")), titolo=str(row.get("Titolo", "")),
                        partita=str(row.get("Partita", "")), rocce=float(row.get("Rocche", 0) or 0),
                        peso=float(row.get("Peso", 0) or 0), label=str(row.get("تحضير خام", "تحضير خام")),
                    )
                    for row in rows
                ])
                lines.append(f"\nRaw yarn to prepare ({len(rows)} row(s)):\n{filato_path}")
            except Exception as exc:
                self._logger.exception("Shared Excel Biglietti (Filato X Tinturia) failed")
                had_error = True
                lines.append(f"\nFilato X Tinturia FAILED:\n{exc}")

        moved = 0
        for partita_col in pg_x_partita_cols or []:
            try:
                moved += move_pg_x_to_orders(self.shared_excel_path, partita_col)
            except Exception as exc:
                self._logger.exception("Shared Excel Biglietti (move to Orders) failed for %s", partita_col)
                had_error = True
                lines.append(f"\nCould not move {partita_col} to Orders:\n{exc}")
        if moved:
            lines.append(f"\nMoved {moved} PG-X row(s) to Orders.")

        if had_error:
            self._set_status("Completed with errors -- see details.")
            self.after(0, lambda: messagebox.showwarning("Biglietti", "\n".join(lines)))
        else:
            self.after(0, lambda: messagebox.showinfo("Biglietti", "\n".join(lines)))
        self.after(0, lambda: self.convert_btn.config(state="normal"))
    def _upload_pgx_gg_file(self, parent_window, datasets, refresh_records, rebuild, partita_gg_var):
        """Read a two-column Partita Col/Partita GG file and assign every PG-X row."""
        if getattr(self, "_pgx_upload_running", False):
            return
        path = filedialog.askopenfilename(parent=parent_window, title="Upload PG-X Partita file", filetypes=[("Excel files", "*.xlsx *.xlsm")])
        if not path:
            return
        import openpyxl
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            ws = wb.active
            headers = [str(cell.value or "").strip().casefold().replace(" ", "") for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            col_col = next((i for i, value in enumerate(headers) if value in {"partitacol", "partita-col", "partita"}), None)
            col_gg = next((i for i, value in enumerate(headers) if value in {"partitagg", "partita-gg", "rawbatch"}), None)
            if col_col is None or col_gg is None:
                wb.close()
                return messagebox.showerror("PG-X Upload", "The file must contain exactly the required headers: Partita Col and Partita GG.", parent=parent_window)
            pairs = [
                (row_number, row[col_col], row[col_gg])
                for row_number, row in enumerate(
                    ws.iter_rows(min_row=2, values_only=True), start=2,
                )
                if row[col_col] not in (None, "") and row[col_gg] not in (None, "")
            ]
            wb.close()
        except Exception as exc:
            return messagebox.showerror("PG-X Upload", str(exc), parent=parent_window)
        if not pairs:
            return messagebox.showinfo("PG-X Upload", "No Partita Col/Partita GG rows were found.", parent=parent_window)

        self._pgx_upload_running = True

        def inspect_file():
            try:
                _codes, densita_map, vmm_ratio_map, _prices, summary = self._load_common_sources()
                pgx_records = load_create_excel_records(self.shared_excel_path, sheet_name="PG-X")
                records_by_partita = {}
                for record in pgx_records:
                    key = _upload_partita_key(record.colored_batch)
                    records_by_partita.setdefault(key, []).append(record)
                ready = []
                missing = []
                errors = []
                for source_row, partita_col, partita_gg in pairs:
                    article_rows = records_by_partita.get(_upload_partita_key(partita_col), [])
                    if not article_rows:
                        errors.append(f"{partita_col}: no matching Partita Col remains in PG-X.")
                        continue
                    status = pg_x_batch_stock_status(
                        [record.article for record in article_rows], str(partita_gg), summary,
                    )
                    if status == "available":
                        ready.append((source_row, partita_col, partita_gg, False))
                    elif status == "missing":
                        missing.append((source_row, partita_col, partita_gg))
                    else:
                        required = ", ".join(sorted({
                            raw_articolo_for(record.article)
                            for record in article_rows
                        }))
                        errors.append(
                            f"{partita_col}: Partita GG {partita_gg} exists in Magazino "
                            f"but is not for the required raw article(s) {required}; not loaded."
                        )
                self.after(
                    0,
                    lambda: confirm_missing(
                        ready, missing, errors, densita_map, vmm_ratio_map, summary,
                    ),
                )
            except Exception as exc:
                self.after(0, lambda exc=exc: fail_upload(exc))

        def confirm_missing(ready, missing, errors, densita_map, vmm_ratio_map, summary):
            approved = []
            if missing:
                details = "\n".join(
                    f"Partita Col {part_col} / Partita GG {part_gg}"
                    for _source_row, part_col, part_gg in missing
                )
                if messagebox.askyesno(
                    "PG-X Upload - Raw batch not in Magazino",
                    "These Partita GG values are not present in Magazino. "
                    "Load them anyway despite the warning?\n\n"
                    f"{details}",
                    parent=parent_window,
                ):
                    approved = [
                        (source_row, part_col, part_gg, True)
                        for source_row, part_col, part_gg in missing
                    ]
                else:
                    errors.extend(
                        f"{part_col}: skipped because Partita GG {part_gg} is not in Magazino."
                        for _source_row, part_col, part_gg in missing
                    )
            threading.Thread(
                target=apply_upload,
                args=(ready + approved, errors, densita_map, vmm_ratio_map, summary),
                daemon=True,
            ).start()

        def apply_upload(assignments, errors, densita_map, vmm_ratio_map, summary):
            updated = 0
            source_rows_removed = 0
            successful_source_rows = []
            for source_row, partita_col, partita_gg, allow_missing in assignments:
                try:
                    result = save_pg_x_partita(
                        self.shared_excel_path, str(partita_col), str(partita_gg),
                        densita_map=densita_map, vmm_ratio_map=vmm_ratio_map,
                        magazino_summary=summary,
                        allow_missing_magazino_batch=allow_missing,
                        move_to_orders=False,
                    )
                    updated += int(result.get("updated", 0))
                    successful_source_rows.append(source_row)
                except Exception as exc:
                    errors.append(f"{partita_col}: {exc}")
            if successful_source_rows:
                try:
                    source_rows_removed = remove_uploaded_pgx_rows(
                        path, successful_source_rows,
                    )
                except Exception as exc:
                    errors.append(
                        "Rows were assigned, but could not be removed from the uploaded file: "
                        f"{exc}"
                    )
            try:
                new_datasets = {
                    "Orders": load_create_excel_records(self.shared_excel_path, sheet_name="Orders"),
                    "PG-X": load_create_excel_records(self.shared_excel_path, sheet_name="PG-X"),
                }
                self.after(
                    0,
                    lambda: finish_upload(
                        updated, source_rows_removed, errors, new_datasets,
                    ),
                )
            except Exception as exc:
                self.after(0, lambda exc=exc: fail_upload(exc))

        def finish_upload(updated, source_rows_removed, errors, new_datasets):
            self._pgx_upload_running = False
            refresh_records(new_datasets)
            rebuild()
            summary_text = f"Assigned {updated} row(s)."
            if source_rows_removed:
                summary_text += (
                    f"\nRemoved {source_rows_removed} processed row(s) from the uploaded file."
                )
            if errors:
                summary_text += "\n\nWarnings / rows not loaded:\n" + "\n".join(errors)
            messagebox.showinfo("PG-X Upload", summary_text, parent=parent_window)

        def fail_upload(exc):
            self._pgx_upload_running = False
            messagebox.showerror("PG-X Upload", str(exc), parent=parent_window)

        threading.Thread(target=inspect_file, daemon=True).start()
    def _smart_auto_assign_pg_x(self, parent_window, datasets, refresh_records, rebuild, partita_gg_var):
        """Automatically match available raw yarn in Magazino Filato with PG-X rows,
        display an interactive approval modal, and move matched rows to Orders."""
        pgx_records = datasets.get("PG-X", [])
        if not pgx_records:
            return messagebox.showinfo("Smart Auto-Assign", "There are no PG-X rows in the current order dataset.", parent=parent_window)

        try:
            _codes, _densita, _vmm, _prices, magazino_summary = self._load_common_sources()
        except Exception as exc:
            return messagebox.showerror("Magazino Load Error", f"Could not load Magazino Filato stock:\n{exc}", parent=parent_window)

        if magazino_summary is None or magazino_summary.empty:
            return messagebox.showinfo("No Stock Data", "Magazino Filato stock is empty or not uploaded.", parent=parent_window)

        stock_pool = {}
        stock_by_lotto = {}
        for row in magazino_summary.itertuples(index=False):
            art = str(getattr(row, "articolo", "") or "").strip().upper()
            partita = str(getattr(row, "partita", "") or "").strip()
            lotto = str(getattr(row, "lotto", "") or "").strip()
            rocche = float(getattr(row, "mag_rocche", 0) or 0)
            peso = float(getattr(row, "mag_peso", 0) or 0)
            if art and partita and rocche > 0:
                item = {"partita": partita, "rocche": rocche, "peso": peso, "lotto": lotto, "articolo": art}
                stock_pool.setdefault(art, []).append(item)
                if lotto:
                    stock_by_lotto.setdefault(lotto.casefold(), []).append(item)

        for art in stock_pool:
            stock_pool[art].sort(key=lambda b: b["rocche"])

        suggestions = []
        used_capacity = {}

        for record in pgx_records:
            raw_b = str(record.raw_batch or "").strip().upper().replace(" ", "")
            if raw_b not in {"", "X", "PG-X", "PGX"}:
                continue

            raw_art = str(record.article or "").strip().upper()
            if raw_art.startswith("C"):
                raw_art = raw_articolo_for(raw_art)

            need_rocche = float(record.quantity_cones or 0)
            comment = str(getattr(record, "commento", "") or getattr(record, "additional_raw", "") or "")
            lotto_match = re.search(r"(?i)PG-([^-\s]+)", comment)
            requested_lotto = lotto_match.group(1).strip() if lotto_match else ""
            candidates = stock_by_lotto.get(requested_lotto.casefold(), []) if requested_lotto and requested_lotto.upper() != "X" else []
            # Lotto is authoritative when present. If no matching Lotto exists,
            # use the requested color's raw article as the fallback.
            if not candidates:
                candidates = stock_pool.get(raw_art, [])

            matching_batch = None
            for b in candidates:
                rem = b["rocche"] - used_capacity.get(b["partita"], 0)
                if rem >= need_rocche:
                    matching_batch = b
                    break

            if matching_batch:
                used_capacity[matching_batch["partita"]] = used_capacity.get(matching_batch["partita"], 0) + need_rocche
                suggestions.append({
                    "record": record,
                    "partita_col": str(record.colored_batch or ""),
                    "cliente": str(record.customer_name or ""),
                    "articolo": str(record.article or ""),
                    "rocche": record.quantity_cones,
                    "suggested_partita": matching_batch["partita"],
                    "available_rocche": matching_batch["rocche"],
                })

        if not suggestions:
            return messagebox.showinfo("Smart Auto-Assign", "No matching raw-yarn stock found in Magazino Filato for current PG-X rows.", parent=parent_window)

        preview = tk.Toplevel(parent_window)
        bind_escape_to_close(preview)
        preview.title("⚡ Smart Auto-Assign Matches")
        preview.geometry("860x460")
        preview.minsize(700, 350)
        preview.transient(parent_window)
        preview.grab_set()

        header_frame = ttk.Frame(preview, padding=10)
        header_frame.pack(fill="x")
        ttk.Label(
            header_frame,
            text=f"Found {len(suggestions)} auto-match(es) for PG-X rows in Magazino Filato!\nReview and select matches to approve:",
            font=("Segoe UI", 10, "bold"), foreground="#16324f",
        ).pack(anchor="w")

        table_frame = ttk.Frame(preview, padding=(10, 0, 10, 10))
        table_frame.pack(fill="both", expand=True)

        columns = ("check", "partita_col", "cliente", "articolo", "rocche", "suggested_partita", "available_rocche")
        headings = {
            "check": "Assign", "partita_col": "Partita Col", "cliente": "Cliente",
            "articolo": "Articolo", "rocche": "Rocche Needed",
            "suggested_partita": "Suggested Partita GG", "available_rocche": "Stock Rocche",
        }
        widths = {"check": 65, "partita_col": 110, "cliente": 160, "articolo": 110, "rocche": 110, "suggested_partita": 140, "available_rocche": 110}

        tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=10)
        for col in columns:
            tree.heading(col, text=headings[col])
            tree.column(col, width=widths[col], anchor="center")
        vsb = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)

        checked = {i: True for i in range(len(suggestions))}

        def render_tree():
            tree.delete(*tree.get_children())
            for idx, item in enumerate(suggestions):
                chk = "☑" if checked[idx] else "☐"
                values = (chk, item["partita_col"], item["cliente"], item["articolo"], item["rocche"], item["suggested_partita"], f"{item['available_rocche']:g}")
                tree.insert("", "end", iid=str(idx), values=values)

        render_tree()

        def toggle_check(event):
            iid = tree.identify_row(event.y)
            if iid:
                idx = int(iid)
                checked[idx] = not checked[idx]
                render_tree()

        tree.bind("<ButtonRelease-1>", toggle_check)

        btn_frame = ttk.Frame(preview, padding=10)
        btn_frame.pack(fill="x")

        def approve_matches():
            selected_suggestions = [suggestions[i] for i, is_chk in checked.items() if is_chk]
            if not selected_suggestions:
                return messagebox.showwarning("No Matches Selected", "Select at least one match to approve.", parent=preview)
            preview.destroy()
            self._apply_auto_assigned_matches(parent_window, selected_suggestions, datasets, refresh_records, rebuild, partita_gg_var)

        ttk.Button(btn_frame, text=f"✅ Approve & Assign Matches ({len(suggestions)})", command=approve_matches).pack(side="right", padx=4)
        ttk.Button(btn_frame, text="Cancel", command=preview.destroy).pack(side="right", padx=4)
    def _apply_auto_assigned_matches(self, parent_window, selected_suggestions, datasets, refresh_records, rebuild, partita_gg_var):
        loading = tk.Toplevel(parent_window)
        loading.title("Assigning Stock...")
        loading.geometry("340x100")
        keep_window_on_top(loading)
        ttk.Label(loading, text="Assigning Partita GG & moving rows to Orders...", padding=20).pack()

        def worker():
            try:
                _codes, densita_map, vmm_ratio_map, _prices, magazino_summary = self._load_common_sources()
                updated_count = 0
                for item in selected_suggestions:
                    partita_col = item["partita_col"]
                    partita_gg = item["suggested_partita"]
                    res = save_pg_x_partita(
                        self.shared_excel_path, partita_col, partita_gg,
                        densita_map=densita_map, vmm_ratio_map=vmm_ratio_map,
                        magazino_summary=magazino_summary,
                        allow_article_mismatch=False,
                    )
                    updated_count += res.get("updated", 0)

                new_datasets = {
                    "Orders": load_create_excel_records(self.shared_excel_path, sheet_name="Orders"),
                    "PG-X": load_create_excel_records(self.shared_excel_path, sheet_name="PG-X"),
                }
                self.after(0, lambda: finish(loading, new_datasets, updated_count))
            except Exception as exc:
                self.after(0, lambda exc=exc: fail(loading, exc))

        def finish(loading, new_datasets, updated_count):
            if loading.winfo_exists():
                loading.destroy()
            refresh_records(new_datasets)
            partita_gg_var.set("")
            rebuild()
            messagebox.showinfo(
                "Auto-Assign Complete",
                f"Assigned Partita GG for {updated_count} row(s) -- they've moved out of PG-X into Orders.\n\n"
                "To print their Biglietti (and export Filato X Tinturia for them), switch to the Orders tab, "
                "check them, and use \"Print Selected Biglietti\" -- they're no longer in PG-X, so "
                "\"Print Assigned PG-X\" won't include them.",
                parent=parent_window,
            )

        def fail(loading, exc):
            if loading.winfo_exists():
                loading.destroy()
            text = str(exc)
            if self._on_notification and ("does not belong to article" in text or "open or locked" in text or "Cannot save" in text):
                self._on_notification("pgx-save-error", "PG-X assignment needs attention", text, "Create (EXCEL+Biglietti)", "high")
            messagebox.showerror("Auto-Assign Error", str(exc), parent=parent_window)

        threading.Thread(target=worker, daemon=True).start()
    def _assigned_pgx_records(self, current_sheet_name: str, record_by_iid: dict[str, object], sheet_by_iid: dict[str, str] = None):
        if current_sheet_name != "PG-X":
            return []
        return [
            record for iid, record in record_by_iid.items()
            if (sheet_by_iid.get(iid) == "PG-X" if sheet_by_iid else iid.startswith("PG-X-"))
            and str(getattr(record, "raw_batch", "") or "").strip().upper().replace(" ", "") not in {"", "X", "PG-X", "PGX"}
        ]
    def _print_assigned_pgx(self, window, current_sheet_name: str, record_by_iid: dict[str, object], sheet_by_iid: dict[str, str] = None):
        if current_sheet_name != "PG-X":
            return
        records = [
            record for iid, record in record_by_iid.items()
            if (sheet_by_iid.get(iid) == "PG-X" if sheet_by_iid else iid.startswith("PG-X-"))
            and str(getattr(record, "raw_batch", "") or "").strip().upper().replace(" ", "") not in {"", "X", "PG-X", "PGX"}
        ]
        if not records:
            return messagebox.showinfo("Print Assigned PG-X", "No PG-X row has a Partita GG assigned yet.", parent=window)
        by_iid = {str(index): record for index, record in enumerate(records)}
        sel = {str(index): True for index in range(len(records))}
        self._print_selected_shared(window, sel, by_iid, move_assigned_pg_x=True)
