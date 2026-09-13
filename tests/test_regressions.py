import unittest
from pathlib import Path
from datetime import date
import tempfile
import pandas as pd
from unittest.mock import patch

from calculate.abbina_calculator import _smallest_fitting_machine
from calculate.abbina_suggestions import titles_compatible
from exporters.biglietti_exporter import _machine_for_count, _filato_rows
from exporters.biglietti_exporter import _get, _read_sheet_rows
from pipelines.master_import import find_files_in_directory
import pipelines.master_import as master_import
from calculate.situazione import compute_delivery_date, compute_delivery_dates
from calculate.lotti import load_lotti
from gui.tabs.overview_tab import _format_display_dates
from gui.tabs.overview_tab import _write_typed_excel_table
import openpyxl
import utility.path_manager as path_manager
import utility.situazione_db as db
from calculate.reports import compute_on_time_delivery, format_partita_timeline
from calculate.prezzi import detect_price_anomalies


class PlanningRegressionTests(unittest.TestCase):
    def test_empty_machine_group_has_no_capacity(self):
        self.assertEqual(_smallest_fitting_machine(0), 0)
        self.assertEqual(_smallest_fitting_machine(-1), 0)

    def test_machine_lookup_rejects_empty_counts(self):
        self.assertEqual(_machine_for_count(None), "")
        self.assertEqual(_machine_for_count(0), "")
        self.assertEqual(_machine_for_count(-2), "")

    def test_special_titles_are_case_and_space_insensitive(self):
        self.assertTrue(titles_compatible(" 30/1 ", "30/1"))
        self.assertFalse(titles_compatible("30/1", "31/1"))

    def test_elvy_reactive_client_colour_keeps_g_marker(self):
        from parsers.dfm_lookup import _extract_colour_code

        self.assertEqual(_extract_colour_code("R.G.4257"), ("4257", "", "G"))

    def test_elvy_reactive_client_colour_matches_dfm(self):
        from parsers.dfm_lookup import lookup_dfm_color

        entries = [{
            "articolo": "C130026S",
            "coloredfm": "364257",
            "cldescr": "EL-G-425711-DOUBLE REATTIVO",
            "twist": "2",
            "titolo": "80/2",
            "date": "2026-09-09",
        }]
        self.assertEqual(
            lookup_dfm_color(
                "G130026S", "R.G.4257", "80/2", "Reactive", "100% Cotton", entries
            ),
            ("364257", "EL-G-425711-DOUBLE REATTIVO"),
        )

    def test_build_dfm_lookup_tolerates_header_case_and_whitespace_drift(self):
        # An ERP export changing "ARTICOLODFM" to " Articolodfm" (case or
        # stray whitespace) between versions must not break the lookup --
        # only a genuinely missing/renamed column should raise an error.
        from parsers.dfm_lookup import build_dfm_lookup
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "DFM.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append([" Articolodfm", "coloredfm", "CLDESCR ", "descrizarticololi", "DataIns"])
            ws.append(["C130027S", "364257", "EL-44011-DOUBLEYARN", "0070  100.00 2", "01/01/2026"])
            wb.save(path)

            entries = build_dfm_lookup(path, prefix="C130")
            self.assertEqual(len(entries), 1)
            self.assertEqual(entries[0]["articolo"], "C130027S")
            self.assertEqual(entries[0]["coloredfm"], "364257")

    def test_build_dfm_lookup_still_errors_on_genuinely_missing_column(self):
        from parsers.dfm_lookup import build_dfm_lookup
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "DFM.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["ARTICOLODFM", "COLOREDFM", "CLDESCR", "DESCRIZARTICOLOLI"])  # DATAINS missing
            ws.append(["C130027S", "364257", "EL-44011-DOUBLEYARN", "0070  100.00 2"])
            wb.save(path)

            with self.assertRaises(ValueError):
                build_dfm_lookup(path, prefix="C130")

    def test_pdf_parser_keeps_multiline_article_description_fields(self):
        from parsers.pdf_parser import PDFParser, ColumnSlot, _group_into_lines

        def word(text, x0, top, width=12):
            return {"text": text, "x0": x0, "x1": x0 + width, "top": top, "bottom": top + 8}

        lines = _group_into_lines([
            word("10", 10, 100),
            word("800", 70, 100),
            word("R.G.4257", 360, 100, 50),
            word("100%", 70, 120, 28),
            word("Cotton", 105, 120, 38),
            word("Pima", 148, 120, 25),
            word("Blend", 178, 120, 30),
            word("Nm", 214, 120, 16),
            word("135/2", 235, 120, 34),
            word("Ne", 275, 120, 14),
            word("80/2", 294, 120, 28),
            word("Reactive", 328, 120, 45),
            word("Griege", 70, 140, 38),
            word("Lot", 114, 140, 18),
            word(":", 136, 140, 5),
            word("99_LOT", 146, 140, 42),
            word("El.09092026", 192, 140, 60),
        ], 4)
        slots = {
            "pos": ColumnSlot("pos", 0, 35),
            "article_area": ColumnSlot("article_area", 35, 350),
            "colour": ColumnSlot("colour", 350, 410),
        }
        rows, _, _ = PDFParser._lines_to_rows(
            PDFParser.__new__(PDFParser), lines, slots, "1479", "9/9/2026", ""
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].article_no, "800")
        self.assertEqual(rows[0].yarn, "100% Cotton Pima Blend")
        self.assertEqual(rows[0].nm, "135/2")
        self.assertEqual(rows[0].ne, "80/2")
        self.assertEqual(rows[0].dye_type, "Reactive")
        self.assertEqual(rows[0].lot, "99_LOT El.09092026")
        self.assertEqual(rows[0].colour, "R.G.4257")

    def test_pdf_parser_recovers_fields_on_first_row_continuation(self):
        from parsers.pdf_parser import PDFParser, ColumnSlot, _group_into_lines

        def word(text, x0, top, width=12):
            return {"text": text, "x0": x0, "x1": x0 + width, "top": top, "bottom": top + 8}

        lines = _group_into_lines([
            word("10", 10, 100),
            word("672.00", 490, 100, 35),
            word("618.24", 610, 100, 35),
            word("800", 70, 108),
            word("R.G.4257", 360, 108, 50),
            word("100%", 70, 120, 28),
            word("Cotton", 105, 120, 38),
            word("Pima", 148, 120, 25),
            word("Blend", 178, 120, 30),
            word("Nm", 214, 120, 16),
            word("135/2", 235, 120, 34),
            word("Ne", 275, 120, 14),
            word("80/2", 294, 120, 28),
            word("Reactive", 328, 120, 45),
            word("Griege", 70, 140, 38),
            word("Lot", 114, 140, 18),
            word(":", 136, 140, 5),
            word("99_LOT", 146, 140, 42),
            word("El.09092026", 192, 140, 60),
        ], 4)
        slots = {
            "pos": ColumnSlot("pos", 0, 35),
            "article_area": ColumnSlot("article_area", 35, 350),
            "colour": ColumnSlot("colour", 350, 410),
            "qty_cones": ColumnSlot("qty_cones", 480, 550),
            "qty_kg": ColumnSlot("qty_kg", 590, 660),
        }
        rows, _, _ = PDFParser._lines_to_rows(
            PDFParser.__new__(PDFParser), lines, slots, "1479", "9/9/2026", ""
        )

        self.assertEqual(rows[0].article_no, "800")
        self.assertEqual(rows[0].colour, "R.G.4257")
        self.assertEqual(rows[0].quantity_cones, 672.0)
        self.assertEqual(rows[0].quantity_kg, 618.24)

    def test_data_start_includes_words_within_first_row_y_tolerance(self):
        from parsers.pdf_parser import PDFParser

        words = [
            {"text": "Pos.", "x0": 10, "x1": 20, "top": 100, "bottom": 106},
            {"text": "10", "x0": 10, "x1": 18, "top": 120.7, "bottom": 127},
            {"text": "800", "x0": 40, "x1": 52, "top": 120.4, "bottom": 127},
        ]
        start = PDFParser._find_data_start(words, words[0])

        self.assertEqual(start, 120.4)

    def test_shared_path_round_trip(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "Produzione.xlsx"
            path.write_bytes(b"placeholder")
            path_manager.save_source("data_prod", path)
            self.assertEqual(
                path_manager.load_source("produzione")["source_path"], str(path)
            )

    def test_elvy_delivery_rules_skip_friday(self):
        row = {"cliente": "3009", "data": "2026-09-11", "mc": 192, "comment": ""}
        self.assertEqual(compute_delivery_date(row), "2026-09-19")
        row["mc"] = 24
        self.assertEqual(compute_delivery_date(row), "2026-09-26")

    def test_elvy_lab_adds_one_week_and_pgx_waits_for_yarn(self):
        row = {"cliente": "ELVY", "data": "2026-09-07", "mc": 192, "new_comment": "Lab"}
        self.assertEqual(compute_delivery_date(row), "2026-09-21")
        row.update({"comment": "PG-X-123", "new_comment": "Filato"})
        self.assertEqual(compute_delivery_date(row), "Bending for yarn")
        row["raw_yarn_match"] = "G130 / 44"
        row["yarn_arrival_date"] = "2026-09-08"
        self.assertEqual(compute_delivery_date(row), "2026-09-15")

    def test_delivery_dates_leave_non_elvy_blank(self):
        frame = pd.DataFrame([{"cliente": "3004", "data": "2026-09-07", "mc": 24}])
        result = compute_delivery_dates(frame)
        self.assertEqual(result.loc[0, "delivery_date"], "")

    def test_master_import_recognizes_lotto_serial_filename(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "ARKMAG-CON-LOTTO_20260829.xlsx"
            path.write_bytes(b"placeholder")
            self.assertEqual(find_files_in_directory(temp_dir)["lotti"], path)

    def test_master_file_finds_data_ordine_sheet_by_content(self):
        # Data Ordine / Dispo Bagno sheets inside a single bundled master
        # workbook have order-specific names (e.g. "MED-D-505449-2026"),
        # never a matchable candidate like "copertura" -- they must be
        # found by sniffing header content instead, the same way loose
        # files in a folder already are.
        with tempfile.TemporaryDirectory() as temp_dir:
            master_path = Path(temp_dir) / "Master.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Copertura"
            ws.append(["N. Bagno", "Machine", "Batch Start"])
            order_ws = wb.create_sheet("MED-D-505449-2026")
            order_ws.append(["Cliente", "DescrizioneAggiuntivaOrdine", "Altro"])
            order_ws.append(["ACME", "Some note", "x"])
            wb.save(master_path)

            matches = master_import._match_sheets(["Copertura", "MED-D-505449-2026"])
            self.assertNotIn("data_ordine", matches)  # not findable by name

            wb2 = openpyxl.load_workbook(master_path, read_only=True)
            try:
                sniffed = master_import._content_match_ordine_sheets(
                    wb2, wb2.sheetnames, set(matches.values())
                )
            finally:
                wb2.close()
            self.assertEqual(sniffed.get("data_ordine"), "MED-D-505449-2026")

    def test_lotti_loader_reads_all_workbook_sheets(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "LOTTI.xlsx"
            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            worksheet.append(["MAGAZZINO", "ARTICOLO", "PARTITA", "ORDINE", "QESI", "LOTTO"])
            worksheet.append([900910, "G130-1", "P-1", 0, 10, "L-1"])
            workbook.create_sheet("Seconda")
            workbook.save(path)

            frame, errors = load_lotti(path)

            self.assertEqual(errors, [])
            self.assertEqual(frame.loc[0, "lotto"], "L-1")

    def test_overview_import_skips_order_only_sources(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            (folder / "Data Ordine.xlsx").write_bytes(b"placeholder")
            (folder / "Dispo Bagno.xlsx").write_bytes(b"placeholder")

            with patch.object(master_import, "_route") as route:
                loaded, skipped = master_import.import_master_directory(
                    folder,
                    situazione_tab=None,
                    magazino_tab=None,
                    skip_keys={"data_ordine", "dispo_bagno"},
                )

            route.assert_not_called()
            self.assertNotIn("Data Ordine", loaded + skipped)
            self.assertNotIn("Dispo Bagno", loaded + skipped)

    def test_master_directory_routes_biglietti_template(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            template = Path(temp_dir) / "Biglietti.docx"
            template.write_bytes(b"template")

            class Biglietti:
                def __init__(self):
                    self.path = None

                def set_template_path(self, path):
                    self.path = path

            target = Biglietti()
            master_import.import_master_directory(
                temp_dir, situazione_tab=None, magazino_tab=None,
                biglietti_tab=target,
            )

            self.assertEqual(target.path, str(template))

    def test_master_file_reports_unfound_sources_as_skipped(self):
        # Regression guard for the previous silent-drop bug: import_master_file
        # used to build its skipped-list from SHEET_CANDIDATES only, so a
        # source with no sheet-name candidate at all (Data Ordine, Dispo
        # Bagno) simply vanished from both loaded and skipped instead of
        # being reported as missing.
        with tempfile.TemporaryDirectory() as temp_dir:
            master_path = Path(temp_dir) / "Master.xlsx"
            wb = openpyxl.Workbook()
            wb.active.title = "Sheet1"
            wb.save(master_path)
            loaded, skipped = master_import.import_master_file(
                str(master_path), situazione_tab=None, magazino_tab=None
            )
            self.assertIn("Data Ordine", skipped)
            self.assertIn("Dispo Bagno", skipped)

    def test_master_file_keeps_shared_magazino_and_lotti_extracts(self):
        class SharedTab:
            def __init__(self):
                self.magazino_paths = []
                self.lotti_paths = []

            def _on_upload_magazino(self, path, cache_path=None):
                self.magazino_paths.append((path, cache_path))

            def _on_upload_lotti(self, path, cache_path=None):
                self.lotti_paths.append((path, cache_path))

        with tempfile.TemporaryDirectory() as temp_dir:
            master_path = Path(temp_dir) / "Master.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "Magazino"
            workbook["Magazino"].append(["Articolo", "Partita", "Peso"])
            workbook.create_sheet("LOTTI").append(["Partita", "Lotto"])
            workbook.save(master_path)

            shared_tab = SharedTab()
            master_import.import_master_file(
                str(master_path), situazione_tab=None, magazino_tab=shared_tab
            )

            self.assertTrue(shared_tab.magazino_paths)
            self.assertTrue(shared_tab.lotti_paths)
            for path, cache_path in shared_tab.magazino_paths + shared_tab.lotti_paths:
                self.assertEqual(path, cache_path)
                self.assertTrue(Path(path).is_file())

    def test_order_headers_can_start_later_and_use_aliases(self):
        class Sheet:
            def iter_rows(self, values_only=True):
                return iter([
                    ["Export generated", None, None],
                    ["ARTICOLO", "Q.TA", "DATA CONSEGNA"],
                    ["C130", 4, "2026-09-06"],
                ])

        rows = _read_sheet_rows(Sheet())
        self.assertEqual(_get(rows[0], "Articolo"), "C130")
        self.assertEqual(_get(rows[0], "Ordinata"), 4)
        self.assertEqual(_get(rows[0], "Consegna"), "2026-09-06")

    def test_report_dates_display_day_first(self):
        frame = pd.DataFrame([{"delivery_date": "2026-09-06", "data": "2026-12-31"}, {"delivery_date": "Bending for yarn", "data": ""}])
        result = _format_display_dates(frame, ["data", "delivery_date"])
        self.assertEqual(result.loc[0, "delivery_date"], "06/09/2026")
        self.assertEqual(result.loc[0, "data"], "31/12/2026")
        self.assertEqual(result.loc[1, "delivery_date"], "Bending for yarn")

    def test_overview_shared_refresh_helper_is_available(self):
        from gui.tabs.overview_tab import OverviewTab
        self.assertTrue(hasattr(OverviewTab, "_refresh_shared_tabs"))

    def test_canonical_ui_tab_imports(self):
        from gui.tabs.biglietti_tab import BigliettiTab
        from gui.tabs.kamal_tab import KamalTab
        from gui.tabs.magazino_filato_tab import MagazinoFilatoTab
        from gui.tabs.situazione_settimana_tab import SettimanaTab
        from gui.tabs.situazione_tab import SituazioneTab

        self.assertTrue(all((BigliettiTab, KamalTab, MagazinoFilatoTab, SettimanaTab, SituazioneTab)))

    def test_canonical_logic_imports(self):
        # The old root-level *_logic.py compatibility shims (prezzi_logic,
        # lotti_logic, situazione_settimana_logic, magazino_logic) were
        # removed once every caller was migrated to import from calculate.*
        # directly (see ARCHITECTURE.md). This just confirms the canonical
        # modules still expose the expected callables.
        from calculate.lotti import summarize_by_partita
        from calculate.prezzi import build_price_lookup, load_prezzi
        from calculate.magazino import summarize_by_partita as magazino_summary
        from calculate.situazione_settimana import summarize as weekly_summary

        self.assertTrue(callable(build_price_lookup))
        self.assertTrue(callable(summarize_by_partita))
        self.assertTrue(callable(magazino_summary))
        self.assertTrue(callable(weekly_summary))
        self.assertTrue(callable(load_prezzi))

    def test_price_lookup_cache_reuses_same_source(self):
        import exporters.biglietti_exporter as biglietti_exporter
        self.assertTrue(hasattr(biglietti_exporter, "_PREZZO_LOOKUP_CACHE"))
        self.assertTrue(hasattr(biglietti_exporter, "_ARTICOLI_MARCA_CACHE"))

    def test_filato_rows_uses_order_total_rocche_and_scaled_magazino_peso(self):
        # Rocche always comes from the order itself now -- stock was
        # already verified earlier, on the Ordine page, before the order
        # went into the system. Peso is the warehouse's per-cone weight
        # rate scaled to how many cones THIS order needs (4), not the
        # warehouse's own stock quantity (12).
        from types import SimpleNamespace
        record = SimpleNamespace(raw_batch="158694", article="C1300275", title="100/2", quantity_cones=4)
        raw_rows = [{"articolo": "G1300275", "peso": 245.32, "تحضير خام": "تحضير خام"}]
        magazino = pd.DataFrame([{
            "articolo": "G130027S", "partita": "158694.0", "mag_rocche": 12,
            "mag_peso": 240.0,
        }])
        result = _filato_rows([record], raw_rows, magazino)
        self.assertEqual(result[0]["Rocche"], 4)
        self.assertEqual(result[0]["Peso"], 80.0)  # (240.0 / 12) * 4

    def test_filato_rows_falls_back_to_raw_sheet_peso_without_magazino(self):
        from types import SimpleNamespace
        record = SimpleNamespace(raw_batch="158694", article="C1300275", title="100/2", quantity_cones=4)
        raw_rows = [{"articolo": "G1300275", "peso": 245.32, "تحضير خام": "تحضير خام"}]
        result = _filato_rows([record], raw_rows, magazino_summary=None)
        self.assertEqual(result[0]["Rocche"], 4)
        self.assertEqual(result[0]["Peso"], 245.32)

    def test_biglietti_filato_uses_fixed_filename_and_shared_writer(self):
        # Create (EXCEL+Biglietti)'s Filato output must follow the same
        # convention as Ordine Kamal/Ordine ELVY: a fixed "Filato x
        # Tinturia.xlsx" name, written via export_filato_full (clear +
        # rewrite each run), not a per-order-named snapshot file.
        from types import SimpleNamespace
        from pipelines.ordini_elvy import RawYarnMatch, export_filato_full, read_filato_tinturia_sheet

        record = SimpleNamespace(raw_batch="158694", article="C1300275", title="100/2", quantity_cones=4)
        raw_rows = [{"articolo": "G1300275", "peso": 245.32, "تحضير خام": "تحضير خام"}]
        magazino = pd.DataFrame([{
            "articolo": "G130027S", "partita": "158694.0", "mag_rocche": 12, "mag_peso": 245.32,
        }])
        filato_rows = _filato_rows([record], raw_rows, magazino)
        matches = [
            RawYarnMatch(
                articolo=r["Articolo"], titolo=r["Titolo"], partita=r["Partita"],
                rocce=r["Rocche"], peso=r["Peso"], label=r["تحضير خام"],
            )
            for r in filato_rows
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / "Filato x Tinturia.xlsx"
            n = export_filato_full(target, matches)
            self.assertEqual(n, 1)
            self.assertTrue(target.is_file())
            rows = read_filato_tinturia_sheet(target)
            self.assertEqual(rows[0].rocce, 4)  # order total, not Magazino's stock count (12)

    def test_elvy_and_el_kamal_workbooks_embed_filato_sheet_like_med(self):
        # Every client's own extract-excel gets a "Filato x Tinturia" tab,
        # not just MED's -- regression guard for the ELVY/EL KAMAL asymmetry
        # where include_filato was left False while MED already had it True.
        from exporters.biglietti_exporter import OrderRecord, export_workbook

        record = OrderRecord(
            customer_code="3009", customer_name="ELVY WEAVING", article="C130027S",
            description="", additional_raw="", color_code="5305", color_name="EL-281311",
            order_no="7777", order_row="1", colored_batch="322813", raw_batch="158694",
            quantity_cones=32, raw_weight=30, dispo="D-00505450-001", bagno="S940",
            title="100/2",
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            for customer in ("ELVY", "EL_KAMAL"):
                path = Path(temp_dir) / f"{customer}.xlsx"
                export_workbook(path, [record], [], include_filato=True, customer=customer)
                wb = openpyxl.load_workbook(path)
                try:
                    self.assertIn("Filato x Tinturia", wb.sheetnames, customer)
                finally:
                    wb.close()

    def test_filato_extract_replaces_previous_order_content(self):
        # A shared "Filato x Tinturia.xlsx" (e.g. Ordine Kamal and Ordine
        # ELVY both pointed at the same folder) must only ever reflect the
        # MOST RECENT extraction -- an older order's raw yarn shouldn't
        # linger next to a newer order's, since it may already be pulled
        # and irrelevant.
        from pipelines.ordini_elvy import export_filato_full, read_filato_tinturia_sheet, RawYarnMatch
        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / "filato.xlsx"
            older_order = [RawYarnMatch(articolo="G130027S", titolo="100/2", partita="158694", rocce=10, peso=100.0)]
            export_filato_full(target, older_order)

            newer_order = [RawYarnMatch(articolo="G999999S", titolo="60/1", partita="222111", rocce=5, peso=50.0)]
            export_filato_full(target, newer_order)

            rows = read_filato_tinturia_sheet(target)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].partita, "222111")  # only the newer order's row remains

    def test_ordini_full_replaces_previous_order_content(self):
        from pipelines.ordini_elvy import export_ordini_full, OrdiniElvyRow
        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / "ordini.xlsx"
            export_ordini_full(target, [OrdiniElvyRow(articolo_delta="C100", coloredfm=5, quantity_cones=10)])
            export_ordini_full(target, [OrdiniElvyRow(articolo_delta="C200", coloredfm=7, quantity_cones=20)])

            wb = openpyxl.load_workbook(target)
            ws = wb["ORDINE VENDITA EGITTO"]
            values = [ws.cell(row=r, column=c).value for r in range(2, ws.max_row + 1) for c in range(1, ws.max_column + 1)]
            wb.close()
            self.assertNotIn("C100", values)  # the older order's row is gone
            self.assertIn("C200", values)

    def test_abbin_zero_normalizes_to_none_and_never_groups(self):
        from pipelines.ordine_med import OrdineMedRow, _normalize_abbin, compute_mc_and_gruppo

        self.assertIsNone(_normalize_abbin(0))
        self.assertIsNone(_normalize_abbin("0"))
        self.assertIsNone(_normalize_abbin(0.0))
        self.assertIsNone(_normalize_abbin(None))
        self.assertIsNone(_normalize_abbin(""))
        self.assertEqual(_normalize_abbin(5), 5)
        self.assertEqual(_normalize_abbin("7"), "7")

        def record(rocc, abbin):
            return OrdineMedRow(
                riga=1, code_org="", titolo="", descr_col="", articolo="C100",
                colore="", rocc=rocc, abbin=_normalize_abbin(abbin), consegna_input="", pt_grg="",
                pt_med="", polmoni="", cliente_note="", nota_grg="", nota_col="",
                kg_note="", fabb="", prezz_note="",
            )

        # Two rows that both had raw ABBIN 0 must NOT be grouped together --
        # each keeps its own Rocche as its M/C, same as a genuinely blank ABBIN.
        zero_a, zero_b = record(10, 0), record(20, 0)
        # Two rows that share a real, non-zero ABBIN must still be grouped
        # and dyed together, combining their Rocche.
        paired_a, paired_b = record(6, 5), record(6, 5)
        compute_mc_and_gruppo([zero_a, zero_b, paired_a, paired_b])

        self.assertEqual(zero_a.mc, 10)
        self.assertEqual(zero_b.mc, 20)
        self.assertEqual(paired_a.mc, 12)
        self.assertEqual(paired_b.mc, 12)

    def test_med_erp_export_replaces_previous_content(self):
        from pipelines.ordine_med import OrdineMedRow, export_erp_order_workbook

        def record(article):
            return OrdineMedRow(
                riga=1, code_org="", titolo="", descr_col="", articolo=article,
                colore="", rocc=1, abbin="", consegna_input="", pt_grg="",
                pt_med="", polmoni="", cliente_note="", nota_grg="", nota_col="",
                kg_note="", fabb="", prezz_note="",
            )

        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / "EXCEL PER ORDINE VENDITA EGITTO.xlsx"
            export_erp_order_workbook(target, [record("C100")])
            export_erp_order_workbook(target, [record("C200")])

            wb = openpyxl.load_workbook(target, data_only=True)
            ws = wb["Dati sistema (B-N)"]
            values = [cell.value for row in ws.iter_rows(min_row=2) for cell in row]
            wb.close()
            self.assertNotIn("C100", values)
            self.assertIn("C200", values)

    def test_filato_reader_handles_short_excel_rows(self):
        from pipelines.ordini_elvy import read_filato_tinturia_sheet
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "short.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Filato x Tinturia"
            sheet.append(["Articolo", "Titolo", "Partita", "Rocche", "Peso", "تحضير خام"])
            sheet.append(["G130027S", "100/2", "158694", 1108])
            workbook.save(source)
            workbook.close()

            rows = read_filato_tinturia_sheet(source)
            self.assertEqual(rows[0].rocce, 1108)
            self.assertEqual(rows[0].peso, 0)

    def test_med_check_articolo_normalizes_article_prefix_and_colour(self):
        from pipelines.ordine_med import OrdineMedRow, compute_check_articolo
        record = OrdineMedRow(
            riga=1, code_org="C130027S", titolo="", descr_col="", articolo="C130027S",
            colore="324229.0", rocc=1, abbin="", consegna_input=None, pt_grg="",
            pt_med="", polmoni="", cliente_note="", nota_grg="", nota_col="",
            kg_note="", fabb=None, prezz_note=None,
        )
        compute_check_articolo([record], {("G130027S", "324229")})
        self.assertEqual(record.check_articolo, "")

    def test_excel_export_keeps_yarn_waiting_comment(self):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        _write_typed_excel_table(
            worksheet,
            pd.DataFrame([
                {"Delivery Date": "2026-09-08"},
                {"Delivery Date": "Bending for yarn"},
            ]),
        )
        self.assertEqual(worksheet.cell(row=2, column=1).value, "08/09/2026")
        self.assertEqual(worksheet.cell(row=2, column=1).data_type, "s")
        self.assertEqual(worksheet.cell(row=3, column=1).value, "Bending for yarn")

    def test_on_time_delivery_scores_and_ranks_clients(self):
        states = {
            "P1": {"cliente": "ELVY", "partita": "P1", "consegna": "2026-09-01", "data_uscita": "2026-08-30"},
            "P2": {"cliente": "ELVY", "partita": "P2", "consegna": "2026-09-01", "data_uscita": "2026-09-05"},
            "P3": {"cliente": "MED", "partita": "P3", "consegna": "2026-09-01", "data_uscita": "2026-09-10"},
            "P4": {"cliente": "MED", "partita": "P4", "consegna": "2026-09-01", "data_uscita": ""},  # still open
        }
        summary = compute_on_time_delivery(states)
        med = summary[summary["cliente"] == "MED"].iloc[0]
        elvy = summary[summary["cliente"] == "ELVY"].iloc[0]
        self.assertEqual(med["shipped"], 1)  # P4 has no Data Uscita yet -- excluded
        self.assertEqual(med["on_time_pct"], 0.0)
        self.assertEqual(med["avg_delay_days"], 9.0)
        self.assertEqual(elvy["shipped"], 2)
        self.assertEqual(elvy["on_time_pct"], 50.0)
        # Worst on-time % sorts first.
        self.assertEqual(summary.iloc[0]["cliente"], "MED")

    def test_partita_history_records_stage_changes_only(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch.object(db, "DB_PATH", str(Path(temp_dir) / "test.db")):
                db.init_db()
                base = {"partita": "P900", "cliente": "ELVY", "articolo": "G130", "colore": "10",
                        "comment": "C.Q", "bagno": "5", "tinto": "", "data_qualita": "",
                        "data_uscita": "", "consegna": "2026-09-10", "old_comment": "", "new_comment": "C.Q"}
                db.upsert_states([base])
                # Re-uploading identical data must NOT add a second history row.
                db.upsert_states([dict(base)])
                stage_2 = dict(base, tinto="2026-09-01", new_comment="Tinto")
                db.upsert_states([stage_2])
                stage_3 = dict(stage_2, data_uscita="2026-09-05", new_comment="Uscita")
                db.upsert_states([stage_3])

                history = db.get_partita_history("P900")
                self.assertEqual(len(history), 3)  # added once, then 2 real stage changes

                timeline = format_partita_timeline(history)
                self.assertEqual(len(timeline), 3)
                self.assertIn("days_in_qc", timeline.columns)
                self.assertIn("consegna", timeline.columns)
                self.assertIn("ritardo", timeline.columns)
                self.assertTrue((timeline["consegna"] == "2026-09-10").all())
                self.assertIn("First seen", timeline.iloc[0]["event"])
                self.assertIn("Dyed on 2026-09-01", timeline.iloc[1]["event"])
                self.assertIn("Shipped on 2026-09-05", timeline.iloc[2]["event"])

    def test_price_anomaly_flags_large_jump_not_small_one(self):
        df = pd.DataFrame([
            {"CLARTICOLO": "G130", "CLCOLORE": "10", "CLDESCR": "Blue",
             "PREZZOLPZ": 10.00, "_START_DATE": "2026-01-01"},
            {"CLARTICOLO": "G130", "CLCOLORE": "10", "CLDESCR": "Blue",
             "PREZZOLPZ": 15.00, "_START_DATE": "2026-06-01"},  # +50%, flagged
            {"CLARTICOLO": "G170", "CLCOLORE": "20", "CLDESCR": "Red",
             "PREZZOLPZ": 10.00, "_START_DATE": "2026-01-01"},
            {"CLARTICOLO": "G170", "CLCOLORE": "20", "CLDESCR": "Red",
             "PREZZOLPZ": 10.30, "_START_DATE": "2026-06-01"},  # +3%, not flagged
        ])
        anomalies = detect_price_anomalies(df, min_pct_change=10.0)
        self.assertEqual(len(anomalies), 1)
        row = anomalies.iloc[0]
        self.assertEqual(row["CLARTICOLO"], "G130")
        self.assertEqual(row["pct_change"], 50.0)

    def test_compute_delay_days_rules(self):
        from calculate.situazione import compute_delay_days

        # 1. Missing data_qualita -> must be blank
        row_no_qualita = {
            "cliente": "MED", "consegna": "2026-09-10", "data_qualita": "",
            "data_uscita": "2026-09-15",
        }
        self.assertEqual(compute_delay_days(row_no_qualita), "")

        # 2. Missing data_uscita -> must be blank
        row_no_uscita = {
            "cliente": "MED", "consegna": "2026-09-10", "data_qualita": "2026-09-08",
            "data_uscita": "",
        }
        self.assertEqual(compute_delay_days(row_no_uscita), "")

        # 3. Non-Elvy customer: late delivery (5 days late -> "5")
        row_med_late = {
            "cliente": "MED", "consegna": "2026-09-10", "data_qualita": "2026-09-08",
            "data_uscita": "2026-09-15",
        }
        self.assertEqual(compute_delay_days(row_med_late), "5")

        # 4. Non-Elvy customer: early delivery (3 days early -> "-3")
        row_med_early = {
            "cliente": "MED", "consegna": "2026-09-10", "data_qualita": "2026-09-05",
            "data_uscita": "2026-09-07",
        }
        self.assertEqual(compute_delay_days(row_med_early), "-3")

        # 5. Non-Elvy customer: on-time delivery (0 days -> "0")
        row_med_ontime = {
            "cliente": "MED", "consegna": "2026-09-10", "data_qualita": "2026-09-08",
            "data_uscita": "2026-09-10",
        }
        self.assertEqual(compute_delay_days(row_med_ontime), "0")

        # 6. ELVY customer: compares against Delivery Date (late 4 days -> "4")
        row_elvy_late = {
            "cliente": "ELVY", "consegna": "2026-09-01", "delivery_date": "2026-09-10",
            "data_qualita": "2026-09-08", "data_uscita": "2026-09-14",
        }
        self.assertEqual(compute_delay_days(row_elvy_late), "4")

        # 7. ELVY customer: early delivery (early 2 days -> "-2")
        row_elvy_early = {
            "cliente": "3009", "consegna": "2026-09-01", "delivery_date": "2026-09-10",
            "data_qualita": "2026-09-06", "data_uscita": "2026-09-08",
        }
        self.assertEqual(compute_delay_days(row_elvy_early), "-2")

    def test_lazy_module_delegates_to_real_module(self):
        from utility.utils import LazyModule, lazy_call
        import pipelines.ordine_med as real_ordine_med

        proxy = LazyModule("pipelines.ordine_med")
        self.assertTrue(callable(proxy.load_ordine))
        self.assertIs(proxy.load_ordine, real_ordine_med.load_ordine)

        wrapped = lazy_call("pipelines.ordine_med", "load_ordine")
        self.assertTrue(callable(wrapped))

    def test_kamal_tab_lazy_names_resolve_to_real_targets(self):
        # Regression guard for the startup-performance fix: kamal_tab.py's
        # heavy names (pandas/openpyxl/pdfplumber-backed) must still
        # actually delegate to the real functions/classes/constants when
        # called, not just avoid crashing at import time.
        import gui.tabs.kamal_tab as kt
        import pipelines.ordine_kamal as real_ordine_kamal
        import calculate.lotti as real_lotti

        self.assertEqual(kt.KAMAL_ARTICLE_PREFIX, real_ordine_kamal.KAMAL_ARTICLE_PREFIX)
        self.assertTrue(callable(kt.load_dfm_cache))
        self.assertIsInstance(kt.load_dfm_cache(), dict)
        self.assertEqual(kt.lotti_logic.RAW_ARTICOLO_PREFIXES, real_lotti.RAW_ARTICOLO_PREFIXES)
        self.assertTrue(callable(kt.build_ordine_kamal_rows))
        self.assertTrue(callable(kt.export_ordini_full))

    def test_days_in_qc_formats_as_clean_integer_no_trailing_zero(self):
        # compute_situation() derives days_in_qc from a float64 column
        # (dt.days on a Timedelta series with some NaT rows upcasts to
        # float) -- naively fillna("").astype(str) on that turns a real
        # value like 3 into the string "3.0" instead of "3".
        import pandas as pd
        col = pd.Series([3.0, float("nan"), 7.0])
        formatted = col.apply(lambda v: str(int(v)) if pd.notna(v) else "")
        self.assertEqual(formatted.tolist(), ["3", "", "7"])

    def test_delivery_date_highlight_flags_dates_under_threshold(self):
        from exporters.biglietti_exporter import _style_sheet
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Articolo", "Delivery Date"])
        ws.append(["G130", date(2026, 9, 10)])
        ws.append(["G140", "Bending for yarn"])
        _style_sheet(ws, days_until_highlight_columns={"Delivery Date": 4})

        rules = list(ws.conditional_formatting)
        formulas = [rule.sqref for rule in rules]
        self.assertTrue(any(str(f) == "B2:B3" for f in formulas), formulas)
        matching = [r for r in rules if str(r.sqref) == "B2:B3"]
        self.assertEqual(len(matching), 1)
        formula_text = list(matching[0].rules[0].formula)[0]
        self.assertIn("ISNUMBER(B2)", formula_text)
        self.assertIn("TODAY()", formula_text)
        self.assertIn("<4", formula_text)

    def test_delivery_date_highlight_absent_when_column_missing(self):
        # MED/EL_KAMAL sheets don't have a "Delivery Date" column -- passing
        # the threshold unconditionally must simply no-op for them.
        from exporters.biglietti_exporter import _style_sheet
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Articolo", "Consegna"])
        ws.append(["G130", date(2026, 9, 10)])
        _style_sheet(ws, days_until_highlight_columns={"Delivery Date": 4})
        self.assertEqual(len(list(ws.conditional_formatting)), 0)

    def test_email_template_round_trips_through_dict(self):
        from utility.email_compose import EmailTemplate
        original = EmailTemplate(to="a@x.com; b@y.com", cc="c@z.com", subject="Order ready", body="Please find attached.")
        restored = EmailTemplate.from_dict(original.to_dict())
        self.assertEqual(restored, original)
        self.assertFalse(restored.is_blank())

    def test_email_template_blank_detection(self):
        from utility.email_compose import EmailTemplate
        self.assertTrue(EmailTemplate().is_blank())
        self.assertTrue(EmailTemplate.from_dict(None).is_blank())
        self.assertTrue(EmailTemplate.from_dict({}).is_blank())
        self.assertFalse(EmailTemplate(subject="hi").is_blank())

    def test_open_outlook_email_raises_clear_error_without_outlook(self):
        # On this (non-Windows / no-Outlook) test environment, win32com
        # simply isn't importable -- confirm that surfaces as a clean,
        # user-facing RuntimeError rather than an unhandled ImportError.
        from utility.email_compose import EmailTemplate, open_outlook_email
        try:
            import win32com.client  # noqa: F401
            self.skipTest("win32com is available in this environment")
        except ImportError:
            pass
        with self.assertRaises(RuntimeError):
            open_outlook_email(EmailTemplate(to="a@x.com", subject="hi"), [])

    def test_open_outlook_email_prefers_already_running_instance(self):
        # Regression guard for the "Welcome to Outlook" setup-wizard bug:
        # GetActiveObject (attach to Outlook that's already open and signed
        # in) must be tried before Dispatch (which can spin up a fresh,
        # unconfigured instance and land on the account-setup wizard).
        import sys
        import types
        from unittest.mock import MagicMock

        fake_mail = MagicMock()
        fake_outlook = MagicMock()
        fake_outlook.CreateItem.return_value = fake_mail

        fake_client = types.ModuleType("win32com.client")
        fake_client.GetActiveObject = MagicMock(return_value=fake_outlook)
        fake_client.Dispatch = MagicMock(side_effect=AssertionError("Dispatch should not be called when GetActiveObject succeeds"))
        fake_win32com = types.ModuleType("win32com")
        fake_win32com.client = fake_client

        with patch.dict(sys.modules, {"win32com": fake_win32com, "win32com.client": fake_client}):
            from utility.email_compose import EmailTemplate, open_outlook_email
            template = EmailTemplate(to="a@x.com", cc="b@x.com", subject="Order Ready", body="See attached.")
            open_outlook_email(template, [])

        fake_client.GetActiveObject.assert_called_once_with("Outlook.Application")
        fake_outlook.CreateItem.assert_called_once_with(0)
        self.assertEqual(fake_mail.To, "a@x.com")
        self.assertEqual(fake_mail.Subject, "Order Ready")
        fake_mail.Display.assert_called_once()

    def test_open_outlook_email_falls_back_to_dispatch_when_not_running(self):
        import sys
        import types
        from unittest.mock import MagicMock

        fake_mail = MagicMock()
        fake_outlook = MagicMock()
        fake_outlook.CreateItem.return_value = fake_mail

        fake_client = types.ModuleType("win32com.client")
        fake_client.GetActiveObject = MagicMock(side_effect=Exception("no running instance"))
        fake_client.Dispatch = MagicMock(return_value=fake_outlook)
        fake_win32com = types.ModuleType("win32com")
        fake_win32com.client = fake_client

        with patch.dict(sys.modules, {"win32com": fake_win32com, "win32com.client": fake_client}):
            from utility.email_compose import EmailTemplate, open_outlook_email
            open_outlook_email(EmailTemplate(to="a@x.com", subject="hi"), [])

        fake_client.Dispatch.assert_called_once_with("Outlook.Application")
        fake_mail.Display.assert_called_once()


if __name__ == "__main__":
    unittest.main()
