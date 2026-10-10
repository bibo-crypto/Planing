"""Follow-ups to the v1.1.3 hardening: Marca fallback, category cache keys and the SQLite-backed DFM."""
from pathlib import Path

import pandas as pd
import pytest

from calculate import prezzi
from exporters import biglietti_exporter as exporter
from parsers import dfm_lookup


def _rows(article, items):
    return [{"CLARTICOLO": article, "CLCOLORE": c, "LIVELLOLPZ": lv, "PREZZOLPZ": p} for c, lv, p in items]


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    from utility import prezzi_db

    monkeypatch.setattr(prezzi, "_CATEGORY_OVERRIDES_PATH", tmp_path / "settings" / "overrides.json")
    monkeypatch.setattr(prezzi_db, "DB_PATH", tmp_path / "orders.sqlite3")
    monkeypatch.setattr(prezzi, "_PREZZI_CACHE", {})
    monkeypatch.setattr(prezzi, "_ARTICOLI_MARCA_CACHE", None)
    prezzi._load_reference_category_map.cache_clear()
    yield tmp_path
    prezzi._load_reference_category_map.cache_clear()


# ---------------------------------------------------------------- Marca fallback

def test_articles_missing_from_articoli_do_not_share_a_bare_customer_category(isolated, monkeypatch):
    monkeypatch.setattr(prezzi, "_load_articoli_marca_lookup", lambda: {"C010111S": "Cotton"})
    frame = pd.DataFrame(
        _rows("C010111S", [("00101", 1, 4.0)])                       # Articoli knows its Marca
        + _rows("C010999S", [("00101", 1, 9.0), ("00102", 1, 9.5)])  # MED family, but not in Articoli
        + _rows("C011888S", [("00101", 1, 2.0)])                     # another one, unrelated prices
    )

    result = prezzi.enrich_categories(frame)

    categories = result.groupby("CLARTICOLO")["CATEGORY"].first().to_dict()
    assert categories["C010111S"] == "MED - Cotton"
    assert categories["C010999S"] == "" and categories["C011888S"] == ""      # no "MED -" catch-all
    assert result.loc[result["CLARTICOLO"] != "C010111S", "_CATEGORY_REVIEW"].all()   # left for Category Review
    lookup = prezzi.build_price_lookup(result)
    assert not any(key[0] == "__CATEGORY__" and key[2].strip(" -") == "med" for key in lookup)
    # and an unpriced colour of one article no longer inherits a price from the other
    assert exporter._prezzo_match_for("C011888S", "00102", lookup)[:3] == ("", False, "")


def test_marca_fallback_still_groups_articles_that_share_a_marca(isolated, monkeypatch):
    monkeypatch.setattr(prezzi, "_load_articoli_marca_lookup",
                        lambda: {"C170123S": "Cashmere", "C170456S": "Cashmere"})
    frame = pd.DataFrame(
        _rows("C170123S", [("00001", 39, 8.5)]) + _rows("C170456S", [("00001", 39, 8.5), ("00002", 39, 9.0)])
    )

    result = prezzi.enrich_categories(frame)

    assert set(result["CATEGORY"]) == {"EL KAMAL - Cashmere"}
    assert exporter._prezzo_match_for("C170123S", "00002", prezzi.build_price_lookup(result))[:3] == (
        9.0, True, "EL KAMAL - Cashmere")


# ---------------------------------------------------------------- cache keys

def test_listini_cache_notices_a_new_articoli_file(isolated, monkeypatch):
    from utility import articoli_cache

    listini = isolated / "Listini.xlsx"
    frame = pd.DataFrame(_rows("C010111S", [("00101", 1, 4.0)]))
    frame["DESCRIZARTICOLOLI"], frame["CLDESCR"] = "art", "col"
    frame.to_excel(listini, index=False)
    articoli = isolated / "Articoli.xlsx"
    articoli.write_text("v1")
    marca = {"C010111S": "Cotton"}
    monkeypatch.setattr(articoli_cache, "load_articoli_cache", lambda: {"source_path": str(articoli)})
    monkeypatch.setattr(exporter, "load_articoli_marca_lookup", lambda: dict(marca))

    first, errors = prezzi.load_prezzi(listini)
    assert not errors and first["CATEGORY"].tolist() == ["MED - Cotton"]

    marca["C010111S"] = "Lino"                       # a corrected Articoli file is uploaded
    articoli.write_text("version two, longer")
    second, _ = prezzi.load_prezzi(listini)

    assert second["CATEGORY"].tolist() == ["MED - Lino"]      # not the category cached with the old file


def test_category_sources_fingerprint_includes_the_articoli_file(isolated, monkeypatch):
    from utility import articoli_cache

    articoli = isolated / "Articoli.xlsx"
    articoli.write_text("a")
    monkeypatch.setattr(articoli_cache, "load_articoli_cache", lambda: {"source_path": str(articoli)})
    before = prezzi.category_sources_fingerprint()

    articoli.write_text("a longer file")

    assert prezzi.category_sources_fingerprint() != before
    monkeypatch.setattr(articoli_cache, "load_articoli_cache", lambda: {})
    assert prezzi.category_sources_fingerprint()[2] is None


# ---------------------------------------------------------------- SQLite-first DFM

@pytest.fixture
def dfm_session(tmp_path, monkeypatch):
    from utility import situazione_db
    monkeypatch.setattr(dfm_lookup, "APP_DATA_DIR", tmp_path)
    monkeypatch.setattr(situazione_db, "DB_PATH", str(tmp_path / "dfm.sqlite3"))
    dfm_lookup._ACTIVE_DFM_CACHE.clear()
    situazione_db.init_db()
    yield tmp_path
    dfm_lookup._ACTIVE_DFM_CACHE.clear()


def test_uploading_the_dfm_again_discards_entries_built_for_other_prefixes(dfm_session, monkeypatch):
    dfm = dfm_session / "DFM.xlsx"
    dfm.write_text("monday")
    built = {"n": 0}

    def fake_build(path, prefix=dfm_lookup.ELVY_ARTICLE_PREFIX):
        built["n"] += 1
        return [{"articolo": f"{prefix}-from-build-{built['n']}"}]

    monkeypatch.setattr(dfm_lookup, "build_dfm_lookup", fake_build)
    dfm_lookup.save_dfm_cache([{"articolo": "elvy-monday"}], dfm.name, dfm)
    kamal_monday = dfm_lookup.load_dfm_entries_by_prefix("C170")
    assert dfm_lookup.load_dfm_entries_by_prefix("C170") == kamal_monday      # cached while nothing changed
    assert built["n"] == 1

    dfm.write_text("tuesday: same file name and path, new export")
    dfm_lookup.save_dfm_cache([{"articolo": "elvy-tuesday"}], dfm.name, dfm)   # explicit upload again

    kamal_tuesday = dfm_lookup.load_dfm_entries_by_prefix("C170")
    assert kamal_tuesday != kamal_monday and built["n"] == 2


def test_saving_a_non_default_prefix_leaves_the_other_entries_alone(dfm_session):
    dfm = dfm_session / "DFM.xlsx"
    dfm.write_text("x")
    dfm_lookup.save_dfm_cache([{"articolo": "elvy"}], dfm.name, dfm)

    dfm_lookup.save_dfm_cache([{"articolo": "kamal"}], dfm.name, dfm, prefix="C170")

    assert dfm_lookup.load_dfm_cache()["entries"] == [{"articolo": "elvy"}]
    assert dfm_lookup.load_dfm_cache(prefix="C170")["entries"] == [{"articolo": "kamal"}]


def test_dfm_data_survives_restart_without_the_original_workbook(dfm_session):
    assert dfm_lookup.active_dfm_source() is None                    # nothing uploaded yet

    dfm = dfm_session / "DFM.xlsx"
    dfm.write_text("x")
    dfm_lookup.save_dfm_cache([{"articolo": "elvy"}], dfm.name, dfm)
    assert dfm_lookup.active_dfm_source() == dfm

    dfm_lookup._ACTIVE_DFM_CACHE.clear()                             # simulate application restart
    dfm.unlink()
    assert dfm_lookup.load_dfm_cache()["entries"] == [{"articolo": "elvy"}]
    assert dfm_lookup.active_dfm_source() is None                    # raw workbook is optional


def test_ordine_med_uses_persisted_dfm_snapshot_not_raw_upload_path():
    source = (Path(__file__).resolve().parent.parent / "gui" / "tabs" / "ordine_med_tab.py").read_text(encoding="utf-8")
    assert 'get_all_uploads().get("dfm"' not in source
    assert "load_dfm_pairs" in source
    assert "active_dfm_source" not in source


def _dfm_workbook(path, rows):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DFM"
    ws.append(["ARTICOLODFM", "COLOREDFM", "CLDESCR", "DESCRIZARTICOLOLI", "DATAINS"])
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return path


@pytest.fixture
def dfm_db(tmp_path, monkeypatch):
    """A throw-away SQLite file, so the DFM import never touches the real database."""
    from parsers import dfm_lookup
    from utility import situazione_db

    monkeypatch.setattr(situazione_db, "DB_PATH", str(tmp_path / "dfm.sqlite3"))
    monkeypatch.setattr(dfm_lookup, "APP_DATA_DIR", tmp_path)
    dfm_lookup._ACTIVE_DFM_CACHE.clear()
    dfm_lookup._ACTIVE_DFM_PAIRS = None
    situazione_db.init_db()
    yield tmp_path
    dfm_lookup._ACTIVE_DFM_CACHE.clear()
    dfm_lookup._ACTIVE_DFM_PAIRS = None


_DFM_ROWS = [
    ("C130001S", "100", "EL-100", "NE 30/1 COTONE", "01/02/2026"),
    ("C170001S", "200", "KM-200", "NE 28/1 CASHMERE", "02/02/2026"),
    ("C010027S", "1040", "MED-1040", "NE 40/1 COTONE", "03/02/2026"),
    ("C011027S", "1041", "MED-1041", "NE 40/1 COTONE", "04/02/2026"),
    ("C010099S", "7777", "", "NE 20/1", "05/02/2026"),       # no CLDESCR: dyed, but not usable for colour matching
    ("X999001S", "55", "OTHER", "NE 10/1", "06/02/2026"),     # a customer the lookups don't cover
]


def test_dfm_workbook_import_persists_all_supported_prefixes_and_pairs_from_one_read(dfm_db, monkeypatch):
    from parsers import dfm_lookup

    path = _dfm_workbook(dfm_db / "DFM.xlsx", _DFM_ROWS)
    reads = []
    original = dfm_lookup.openpyxl.load_workbook
    monkeypatch.setattr(dfm_lookup.openpyxl, "load_workbook", lambda *a, **k: (reads.append(1), original(*a, **k))[1])

    counts = dfm_lookup.save_dfm_workbook(path)

    assert reads == [1]                                              # one read, not one per prefix
    assert counts == {"C130": 1, "C170": 1, "C010": 1, "C011": 1}   # the CLDESCR-less row is not a lookup row
    assert dfm_lookup.load_dfm_cache("C010")["entries"][0]["articolo"] == "C010027S"
    assert dfm_lookup.load_dfm_cache("C011")["entries"][0]["coloredfm"] == "1041"
    assert dfm_lookup.load_dfm_cache("C130")["entries"][0]["articolo"] == "C130001S"
    assert dfm_lookup.load_dfm_cache("C170")["entries"][0]["articolo"] == "C170001S"


def test_check_articolo_pairs_come_from_sqlite_and_cover_every_dfm_row(dfm_db):
    from parsers import dfm_lookup
    from pipelines.ordine_med import OrdineMedRow, compute_check_articolo

    dfm_lookup.save_dfm_workbook(_dfm_workbook(dfm_db / "DFM.xlsx", _DFM_ROWS))
    dfm_lookup._ACTIVE_DFM_CACHE.clear()          # application restart: memory is gone,
    dfm_lookup._ACTIVE_DFM_PAIRS = None           # the workbook may be gone too
    (dfm_db / "DFM.xlsx").unlink()

    pairs = dfm_lookup.load_dfm_pairs()

    assert {("C130001S", "100"), ("C170001S", "200"), ("C010027S", "1040"), ("C011027S", "1041"),
            ("C010099S", "7777"), ("X999001S", "55")} <= pairs

    def record(articolo, colore):
        return OrdineMedRow(
            riga=1, code_org=articolo, titolo="", descr_col="", articolo=articolo, colore=colore, rocc=1,
            abbin="", consegna_input=None, pt_grg="", pt_med="", polmoni="", cliente_note="", nota_grg="",
            nota_col="", kg_note="", fabb=None, prezz_note=None,
        )

    known_without_description, dyed_med, never_dyed = record("C010099S", "7777"), record("C011027S", "1041"), record("C010027S", "9999")
    compute_check_articolo([known_without_description, dyed_med, never_dyed], pairs)
    assert known_without_description.check_articolo == ""      # dyed before, even though CLDESCR was empty
    assert dyed_med.check_articolo == ""
    assert never_dyed.check_articolo != ""                     # a genuinely new Articolo+Colore


def test_a_new_dfm_upload_replaces_the_pairs_of_the_previous_one(dfm_db):
    from parsers import dfm_lookup

    dfm_lookup.save_dfm_workbook(_dfm_workbook(dfm_db / "DFM.xlsx", _DFM_ROWS))
    dfm_lookup.save_dfm_workbook(_dfm_workbook(dfm_db / "DFM.xlsx", [("C010027S", "1040", "MED-1040", "NE 40/1", "03/02/2026")]))

    assert dfm_lookup.load_dfm_pairs() == {("C010027S", "1040")}
    assert dfm_lookup.load_dfm_cache("C130")["entries"] == []        # no stale Elvy rows from the old file


def test_dfm_without_pairs_snapshot_falls_back_to_the_lookup_rows(dfm_db):
    from parsers import dfm_lookup

    dfm_lookup.save_dfm_cache([{"articolo": "C010027S", "coloredfm": "1040"}], "DFM.xlsx", prefix="C010")
    dfm_lookup._ACTIVE_DFM_PAIRS = None

    assert ("C010027S", "1040") in dfm_lookup.load_dfm_pairs()


def test_med_check_articolo_c010_exact_article_and_colour_is_not_new():
    from pipelines.ordine_med import OrdineMedRow, compute_check_articolo

    record = OrdineMedRow(
        riga=1, code_org="C010027S", titolo="", descr_col="", articolo="C010027S",
        colore="1040", rocc=1, abbin="", consegna_input=None, pt_grg="",
        pt_med="", polmoni="", cliente_note="", nota_grg="", nota_col="",
        kg_note="", fabb=None, prezz_note=None,
    )
    compute_check_articolo([record], {("C010027S", "1040")})
    assert record.check_articolo == ""


def test_med_check_articolo_missing_pair_is_new():
    from pipelines.ordine_med import OrdineMedRow, compute_check_articolo

    record = OrdineMedRow(
        riga=1, code_org="C010027S", titolo="", descr_col="", articolo="C010027S",
        colore="1041", rocc=1, abbin="", consegna_input=None, pt_grg="",
        pt_med="", polmoni="", cliente_note="", nota_grg="", nota_col="",
        kg_note="", fabb=None, prezz_note=None,
    )
    compute_check_articolo([record], {("C010027S", "1040")})
    assert record.check_articolo == "NEW"
